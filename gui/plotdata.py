"""Data behind a matplotlib figure: collect it per panel and export it (clipboard, Excel, CSV).

Lines, scatter points, bars, histograms, error bars and images are read from the figure itself. Plots whose data
cannot be recovered from the drawn objects (contour and surface grids, confidence bands, box plots) register their
data with `register`.
"""
import csv
import os
import re

import numpy as np
from matplotlib.collections import PathCollection
from matplotlib.container import BarContainer, ErrorbarContainer


def register(ax, title, header, rows, only=False):
    """Attach a data table to an axes. With only=True the drawn objects of this axes are not read as well."""
    if not hasattr(ax, "_regen_tables"):
        ax._regen_tables = []
    ax._regen_tables.append((title, list(header), [list(r) for r in rows]))
    if only:
        ax._regen_only = True


def register_columns(ax, title, columns, only=False):
    """register() for equally long columns given as {name: values}."""
    names = list(columns)
    cols = [np.asarray(columns[n]).ravel() for n in names]
    register(ax, title, names, zip(*cols), only)


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return None if not np.isfinite(f) else f


def _panel_name(ax, n):
    t = " ".join((ax.get_title() or ax.get_title("left") or "").split())
    if t:
        return t
    xl = ax.get_xlabel().strip()
    return f"Panel {n} ({xl})" if xl else f"Panel {n}"


def _label(artist, default):
    lb = artist.get_label() if hasattr(artist, "get_label") else ""
    return default if not lb or lb.startswith("_") else " ".join(lb.split())


def _ticks(ax, axis):
    pos = ax.get_xticks() if axis == "x" else ax.get_yticks()
    labels = ax.get_xticklabels() if axis == "x" else ax.get_yticklabels()
    return {round(float(p), 6): t.get_text() for p, t in zip(pos, labels) if t.get_text()}


def _colorbar_label(fig, mappable):
    for a in fig.axes:
        cb = getattr(a, "_colorbar", None)
        if cb is not None and cb.mappable is mappable:
            return cb.ax.get_ylabel() or cb.ax.get_xlabel() or "Color value"
    return "Color value"


def _axes_tables(fig, ax, n):
    name = _panel_name(ax, n)
    xl, yl = ax.get_xlabel() or "x", ax.get_ylabel() or "y"
    is3d = hasattr(ax, "get_zlim")
    zl = ax.get_zlabel() if is3d else ""
    out = []
    skip = set()

    for c in ax.containers:
        if isinstance(c, ErrorbarContainer):
            data_line, caps, bars = c.lines
            skip.add(id(data_line))
            skip.update(id(x) for x in caps)
            rows = []
            ticks = _ticks(ax, "x")
            if bars:
                for seg in bars[0].get_segments():
                    (x0, y0), (x1, y1) = seg[0], seg[-1]
                    rows.append([ticks.get(round(float(x0), 6), _num(x0)), _num(min(y0, y1)), _num(max(y0, y1))])
            if data_line is not None:
                xs, ys = data_line.get_xdata(), data_line.get_ydata()
                for r, y in zip(rows, ys):
                    r.insert(1, _num(y))
                if not rows:
                    rows = [[_num(x), _num(y)] for x, y in zip(xs, ys)]
            head = [xl, yl, "Lower", "Upper"][:len(rows[0])] if rows else [xl, yl]
            out.append((f"{name} - {_label(c, 'Error bars')}", head, rows))
        elif isinstance(c, BarContainer):
            horizontal = getattr(c, "orientation", None) == "horizontal"
            ticks = _ticks(ax, "y" if horizontal else "x")
            rows = []
            for p in c.patches:
                if horizontal:
                    lo, w, val, base = p.get_y(), p.get_height(), p.get_width(), p.get_x()
                else:
                    lo, w, val, base = p.get_x(), p.get_width(), p.get_height(), p.get_y()
                center = lo + w / 2
                rows.append([ticks.get(round(center, 6), ""), _num(lo), _num(lo + w), _num(val), _num(base)])
            pos = yl if horizontal else xl
            val = xl if horizontal else yl
            head = ["Label", f"{pos} (from)", f"{pos} (to)", val, "Base"]
            if not any(r[0] for r in rows):
                head, rows = head[1:], [r[1:] for r in rows]
            if rows and all(r[-1] == 0 for r in rows):
                head, rows = head[:-1], [r[:-1] for r in rows]
            out.append((f"{name} - {_label(c, 'Bars')}", head, rows))

    for i, ln in enumerate(ax.get_lines()):
        if id(ln) in skip:
            continue
        if is3d and hasattr(ln, "get_data_3d"):
            xs, ys, zs = (np.asarray(v) for v in ln.get_data_3d())
            rows = [[_num(a), _num(b), _num(c)] for a, b, c in zip(xs, ys, zs)]
            out.append((f"{name} - {_label(ln, f'Line {i + 1}')}", [xl, yl, zl or "z"], rows))
            continue
        xs, ys = np.asarray(ln.get_xdata(), dtype=object), np.asarray(ln.get_ydata(), dtype=object)
        if len(xs) == 0:
            continue
        lb = _label(ln, f"Line {i + 1}")
        tr = ln.get_transform()
        if tr == ax.get_yaxis_transform() and len(set(map(str, ys))) == 1:      # horizontal reference line
            out.append((f"{name} - {lb}", ["Horizontal line at", yl], [["", _num(ys[0])]]))
            continue
        if tr == ax.get_xaxis_transform() and len(set(map(str, xs))) == 1:      # vertical reference line
            out.append((f"{name} - {lb}", ["Vertical line at", xl], [["", _num(xs[0])]]))
            continue
        out.append((f"{name} - {lb}", [xl, yl], [[_num(a), _num(b)] for a, b in zip(xs, ys)]))

    for i, coll in enumerate(ax.collections):
        if not isinstance(coll, PathCollection):
            continue
        lb = _label(coll, f"Points {i + 1}")
        if is3d and hasattr(coll, "_offsets3d"):
            xs, ys, zs = (np.asarray(v, float).ravel() for v in coll._offsets3d)
            out.append((f"{name} - {lb}", [xl, yl, zl or "z"],
                        [[_num(a), _num(b), _num(c)] for a, b, c in zip(xs, ys, zs)]))
            continue
        off = np.asarray(coll.get_offsets(), float)
        if off.size == 0:
            continue
        cols, head = [off[:, 0], off[:, 1]], [xl, yl]
        arr = coll.get_array()
        if arr is not None and len(np.ravel(arr)) == len(off):
            cols.append(np.ravel(arr))
            head.append(_colorbar_label(fig, coll))
        sizes = coll.get_sizes()
        if len(sizes) == len(off) and len(np.unique(sizes)) > 1:
            cols.append(sizes)
            head.append("Marker size")
        out.append((f"{name} - {lb}", head, [[_num(v) for v in r] for r in zip(*cols)]))

    for i, im in enumerate(ax.get_images()):
        A = np.asarray(im.get_array(), float)
        if A.ndim != 2:
            continue
        xt = [t.get_text() for t in ax.get_xticklabels()]
        yt = [t.get_text() for t in ax.get_yticklabels()]
        cols = xt if len(xt) == A.shape[1] and all(xt) else [str(c + 1) for c in range(A.shape[1])]
        rows_l = yt if len(yt) == A.shape[0] and all(yt) else [str(r + 1) for r in range(A.shape[0])]
        out.append((f"{name} - Matrix {i + 1}", [""] + cols,
                    [[rl] + [_num(v) for v in A[r]] for r, rl in enumerate(rows_l)]))
    return out


def figure_tables(fig):
    """List of (title, header, rows) for every panel of the figure (colorbars are skipped)."""
    tables = []
    n = 0
    for ax in fig.axes:
        if getattr(ax, "_colorbar", None) is not None or not ax.get_visible():
            continue
        n += 1
        for title, head, rows in getattr(ax, "_regen_tables", []):
            tables.append((f"{_panel_name(ax, n)} - {title}" if title else _panel_name(ax, n), head, rows))
        if not getattr(ax, "_regen_only", False):
            tables.extend(t for t in _axes_tables(fig, ax, n) if t[2])
    return [(" ".join(t.split()), [" ".join(str(h).split()) for h in head], rows) for t, head, rows in tables]


def tables_to_rows(tables):
    """Tables stacked below each other (title row, header row, data, empty row), e.g. for the clipboard."""
    rows = []
    for title, head, data in tables:
        rows.append([title])
        rows.append(head)
        rows.extend(["" if v is None else (repr(v) if isinstance(v, float) else str(v)) for v in r] for r in data)
        rows.append([])
    return rows


def _sheet_name(n, title, used):
    """Short, unique sheet name: running number and the series part of the title."""
    part = title.split(" - ")[-1] if " - " in title else title
    base = f"{n} " + re.sub(r"[\[\]:*?/\\]", " ", part).strip()
    name, k = base[:31], 2
    while name.lower() in used:
        name, k = f"{base[:26]} ({k})", k + 1
    used.add(name.lower())
    return name


def write_tables(path, tables):
    """Write tables to .xlsx (one sheet per table) or .csv (tables below each other)."""
    if os.path.splitext(path)[1].lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            for title, head, data in tables:
                w.writerow([title])
                w.writerow(head)
                w.writerows([["" if v is None else v for v in r] for r in data])
                w.writerow([])
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for n, (title, head, data) in enumerate(tables or [("Data", [], [])], 1):
        ws = wb.create_sheet(_sheet_name(n, title, used))
        ws.append([title])
        ws["A1"].font = Font(bold=True)
        ws.append(head)
        for c in ws[2]:
            c.font = Font(bold=True)
        for r in data:
            ws.append(r)
        for col in ws.iter_cols(min_row=2, max_row=2):
            ws.column_dimensions[col[0].column_letter].width = max(12, min(40, len(str(col[0].value or "")) + 2))
    wb.save(path)


def export_dialog(parent, tables, default_name="data"):
    """Ask for a file name and write the tables. Returns the path or None."""
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    if not tables:
        QMessageBox.information(parent, "Export data", "There are no data to export.")
        return None
    name = re.sub(r"[^\w\- ]+", "", default_name).strip() or "data"
    path, flt = QFileDialog.getSaveFileName(parent, "Export data", name + ".xlsx",
                                            "Excel workbook (*.xlsx);;CSV file (*.csv)")
    if not path:
        return None
    if not os.path.splitext(path)[1]:
        path += ".csv" if "csv" in flt.lower() else ".xlsx"
    try:
        write_tables(path, tables)
    except OSError as exc:
        QMessageBox.warning(parent, "Export data", f"The file could not be written:\n{exc}")
        return None
    return path
