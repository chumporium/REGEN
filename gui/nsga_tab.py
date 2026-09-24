"""NSGA-II / NSGA-III tab: multi-objective optimization, Pareto front, compromise solution selection (TOPSIS / knee),
convergence monitoring (hypervolume, automatic stop), population adequacy check, and ANN vs RSM front comparison."""
import numpy as np
from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QBrush, QColor
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDoubleSpinBox, QFileDialog, QFormLayout, QGroupBox,
                               QHBoxLayout, QHeaderView, QLabel, QMessageBox, QProgressBar, QPushButton, QSpinBox,
                               QSplitter, QTableWidget, QTableWidgetItem, QTabWidget, QVBoxLayout, QWidget)

from doe import models, nsga2

from . import plots
from .common import REPORT_CSS, MplCanvas, Task, fmt, model_source_combo, parse_float
from .copying import ReportBrowser, install_table_copy

GOALS = [("none", "Ignore"), ("maximize", "Maximize"), ("minimize", "Minimize"), ("target", "Target")]
COLS = ["Objective", "Model", "Goal", "Target", "Lower limit (constraint)", "Upper limit (constraint)",
        "TOPSIS weight"]
ALG_NAME = {"nsga2": "NSGA-II", "nsga3": "NSGA-III"}
C_RSM, C_ANN = "#1f5f99", "#d9480f"
KINDS = [("front", "Pareto front (2D / 3D / matrix)"), ("front2d", "Pareto 2D (choose axes)"),
         ("matrix", "Objective matrix (each pair)"), ("parallel", "Parallel coordinates"),
         ("convergence", "Convergence (hypervolume)")]
CMP_KINDS = [("cmp_front", "ANN vs RSM: 2D front"), ("cmp_cross", "ANN vs RSM: cross-prediction"),
             ("cmp_factors", "ANN vs RSM: factors")]


def objective_label(project, o):
    base = project.responses[o["index"]].name
    return f"POE {base}" if o.get("poe") else base


def _goal_word(o):
    return {"maximize": "max", "minimize": "min", "target": f"target {fmt(o.get('target'))}"}.get(o["goal"], "")


def to_F(o, y):
    """Response value -> objective value to be minimized."""
    if o["goal"] == "maximize":
        return -y
    if o["goal"] == "target":
        return np.abs(y - o["target"])
    return y


def objective_predict(project, o, X, source):
    m = project._model_for(o["index"], source)
    if o.get("poe"):
        m = models.PoeModel(m, project.sd_coded())
    return m.predict(X)


def pair_front(Fa, Fb):
    """Indices of the non-dominated points when only these two objectives are considered (sorted by Fa)."""
    o = np.lexsort((Fb, Fa))
    keep, best = [], np.inf
    for i in o:
        if Fb[i] < best - 1e-12:
            keep.append(i)
            best = Fb[i]
    return np.array(keep, int)


def _names(project, res):
    return [objective_label(project, o) for o in res["active"]]


def _axis_label(project, res, k):
    o = res["active"][k]
    return f"{objective_label(project, o)} ({_goal_word(o)})"


def _title(res, extra=""):
    alg = ALG_NAME.get(res.get("algorithm"), "NSGA-II")
    src = {"rsm": " - all RSM", "ann": " - all ANN"}.get(res.get("source"), "")
    return f"Pareto {alg}{src}: {len(res['X'])} solutions{extra}"


# ================================================================ graphs
def draw_pareto(fig, project, res, highlight=None, kind="front", opts=None):
    opts = opts or {}
    act = res["active"]
    m = len(act)
    Y = np.column_stack(res["Yobj"])
    names = _names(project, res)
    it, ik = res["i_topsis"], res["i_knee"]
    if kind == "front" and m == 2:
        kind, opts = "front2d", {"x": 0, "y": 1, "color": "topsis"}
    if kind == "front2d":
        _draw_front2d(fig, project, res, highlight, opts)
    elif kind == "front" and m == 3:
        ax = fig.add_subplot(111, projection="3d")
        sc = ax.scatter(Y[:, 0], Y[:, 1], Y[:, 2], c=res["topsis"], cmap="viridis", s=30, depthshade=False)
        ax.scatter(*Y[it], marker="*", s=300, color="#ffd400", edgecolor="k", depthshade=False, label="TOPSIS")
        ax.scatter(*Y[ik], marker="D", s=90, color="#e8590c", edgecolor="k", depthshade=False, label="Knee")
        if highlight is not None:
            ax.scatter(*Y[highlight], s=260, facecolor="none", edgecolor="#b42318", lw=2, depthshade=False,
                       label="Selected")
        fig.colorbar(sc, ax=ax, shrink=0.6, label="TOPSIS score")
        ax.set_xlabel(names[0])
        ax.set_ylabel(names[1])
        ax.set_zlabel(names[2])
        ax.set_title(_title(res, " (3 objectives)"))
        ax.legend()
    elif kind == "convergence":
        _draw_convergence(fig, res)
    elif kind in ("front", "matrix"):
        axs = fig.subplots(m, m, squeeze=False)
        for a in range(m):
            for b in range(m):
                ax = axs[a, b]
                if a == b:
                    ax.hist(Y[:, a], bins=15, color=plots.BLUE, alpha=0.7)
                else:
                    ax.scatter(Y[:, b], Y[:, a], c=res["topsis"], cmap="viridis", s=10)
                    ax.scatter(Y[it, b], Y[it, a], marker="*", s=120, color="#ffd400", edgecolor="k", zorder=5)
                    if highlight is not None:
                        ax.scatter(Y[highlight, b], Y[highlight, a], s=90, facecolor="none", edgecolor="#b42318",
                                   lw=1.5, zorder=6)
                ax.tick_params(labelsize=6)
                if a == m - 1:
                    ax.set_xlabel(_axis_label(project, res, b), fontsize=8)
                if b == 0:
                    ax.set_ylabel(_axis_label(project, res, a), fontsize=8)
        fig.suptitle(_title(res, ": every pair of objectives (star = TOPSIS choice)"), fontsize=10)
    else:
        _draw_parallel(fig, project, res, highlight)


def _draw_front2d(fig, project, res, highlight, opts):
    """2D front projection: two objectives as axes, the other objectives as point color and size."""
    act = res["active"]
    m = len(act)
    Y = np.column_stack(res["Yobj"])
    F = res["F"]
    a, b = int(opts.get("x", 0)), int(opts.get("y", 1 if m > 1 else 0))
    col, size = opts.get("color", "topsis"), opts.get("size")
    ax = fig.add_subplot(111)
    if isinstance(col, int) and 0 <= col < m:
        cval, clabel = Y[:, col], _axis_label(project, res, col)
    else:
        cval, clabel = res["topsis"], "TOPSIS score"
    if isinstance(size, int) and 0 <= size < m:
        z = F[:, size]
        g = 1 - (z - z.min()) / (np.ptp(z) or 1)          # better = larger
        sz = 14 + 150 * g
    else:
        sz = np.full(len(Y), 40.0)
    pf = pair_front(F[:, a], F[:, b]) if a != b else np.array([], int)
    if len(pf) > 1:
        ax.plot(Y[pf, a], Y[pf, b], "-", color="#9fb4cf", lw=1.2, zorder=1,
                label=f"Front of these 2 objectives ({len(pf)} points)")
    sc = ax.scatter(Y[:, a], Y[:, b], c=cval, cmap="viridis", s=sz, edgecolor="k", lw=0.4, zorder=3)
    fig.colorbar(sc, ax=ax, label=clabel)
    it, ik = res["i_topsis"], res["i_knee"]
    ax.scatter(Y[it, a], Y[it, b], marker="*", s=380, color="#ffd400", edgecolor="k", zorder=5, label="TOPSIS")
    ax.scatter(Y[ik, a], Y[ik, b], marker="D", s=110, color="#e8590c", edgecolor="k", zorder=4, label="Knee")
    if highlight is not None:
        ax.scatter(Y[highlight, a], Y[highlight, b], s=260, facecolor="none", edgecolor="#b42318", lw=2,
                   zorder=6, label="Selected")
    ax.set_xlabel(_axis_label(project, res, a))
    ax.set_ylabel(_axis_label(project, res, b))
    title = _title(res)
    if m > 2:
        title += f"\n2D projection of {m} objectives"
        if isinstance(size, int) and 0 <= size < m:
            title += f"; point size: {_names(project, res)[size]} (larger = better)"
        title += ".\nPoints that look worse in this plane are better in other objectives."
    ax.set_title(title, fontsize=9)
    ax.grid(alpha=0.3)
    ax.legend(loc="best", fontsize=8)


def _draw_convergence(fig, res):
    hist = res["history"]
    g = [h["gen"] for h in hist]
    ax1, ax2 = fig.subplots(2, 1, sharex=True)
    hv = res.get("hv_curve") or []
    if hv:
        ax1.plot([p[0] for p in hv], [100 * p[1] for p in hv], color=plots.BLUE, lw=2, label="Hypervolume")
    ax1.set_ylabel("Hypervolume (% of max)")
    ax1b = ax1.twinx()
    ax1b.plot(g, [h["n_front"] for h in hist], color="#8a8a8a", lw=1, ls="--", label="Solutions in first front")
    ax1b.set_ylabel("Solutions in front", color="#666")
    chk = res.get("check") or {}
    ca = chk.get("converged_at")
    for ax in (ax1, ax2):
        if ca:
            ax.axvline(ca, color="#1a7f37", lw=1.2, ls=":")
        if res.get("stop_gen"):
            ax.axvline(res["stop_gen"], color="#b42318", lw=1.2)
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax1b.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="lower right", fontsize=8)
    ax1.grid(alpha=0.3)
    title = f"Convergence {ALG_NAME.get(res.get('algorithm'), 'NSGA-II')}: {res.get('gens_run', len(hist))} generations"
    if res.get("stop_gen") and "converged" in (res.get("stop_reason") or ""):
        title += f", automatic stop at generation {res['stop_gen']}"
    ax1.set_title(title, fontsize=10)
    w = res.get("stop_window", 20)
    chk_h = [h for h in hist if h.get("conv")]
    if chk_h:
        gx = np.array([h["gen"] for h in chk_h])
        gain = np.array([h["gain"] for h in chk_h]) * 100
        thr = np.array([nsga2.stop_tol_eff(h, res.get("stop_tol", 1e-3)) for h in chk_h]) * 100
        shift = np.array([h["shift"] for h in chk_h]) * 100
        ax2.semilogy(gx, np.maximum(np.abs(gain), 1e-4), color=plots.BLUE, lw=1.5,
                     label=f"Mean hypervolume gain (last {w} gen. vs previous, %)")
        neg = gain < 0
        if neg.any():
            ax2.scatter(gx[neg], np.maximum(np.abs(gain[neg]), 1e-4), s=10, color=plots.BLUE, marker="v",
                        label="decrease (no progress)")
        ax2.step(gx, thr, where="mid", color=plots.BLUE, ls="--", lw=0.8, label="Threshold (tolerance / noise level)")
        ax2.semilogy(gx, np.maximum(shift, 1e-4), color="#e8590c", lw=1.2, label="Ideal point shift (% of range)")
        ax2.axhline(1.0, color="#e8590c", ls="--", lw=0.8)
        ax2.legend(fontsize=7, loc="upper right")
    else:
        ax2.text(0.5, 0.5, f"At least {2 * w} generations are needed to measure the change", ha="center", va="center",
                 transform=ax2.transAxes, color="#555")
    ax2.set_xlabel("Generation" + ("   (dotted green line: convergence criterion first met; red line: "
                                   "stopped)" if ca or res.get("stop_gen") else ""), fontsize=8)
    ax2.set_ylabel("Change (%)")
    ax2.grid(alpha=0.3, which="both")


def _factor_axes(project, Xs):
    """Scale factors to the design range [0, 1] for parallel coordinates."""
    lo_c, hi_c = project.coded_bounds()
    lo_a, hi_a = project.to_actual(lo_c[None])[0], project.to_actual(hi_c[None])[0]
    out = [[] for _ in Xs]
    top, bottom, labels = [], [], []
    for i, f in enumerate(project.factors):
        a, b = sorted((lo_a[i], hi_a[i]))
        for n, X in enumerate(Xs):
            out[n].append((X[:, i] - a) / (b - a) if b > a else np.full(len(X), 0.5))
        top.append(project.format_value(i, b, 4))
        bottom.append(project.format_value(i, a, 4))
        labels.append(f.name)
    return out, top, bottom, labels


def _draw_parallel(fig, project, res, highlight):
    ax = fig.add_subplot(111)
    act = res["active"]
    Y = np.column_stack(res["Yobj"])
    names = _names(project, res)
    it = res["i_topsis"]
    (Z,), top, bottom, labels = _factor_axes(project, [project.to_actual(res["X"])])
    for n, o in enumerate(act):
        v = Y[:, n]
        lo, hi = v.min(), v.max()
        z = (v - lo) / (hi - lo) if hi > lo else np.full(len(v), 0.5)
        if o["goal"] == "minimize":        # top = better for all objectives
            z, lo, hi = 1 - z, hi, lo
        elif o["goal"] == "target":
            d = np.abs(v - o["target"])
            z = 1 - (d - d.min()) / (np.ptp(d) or 1)
            lo, hi = "far", f"= {fmt(o['target'], 4)}"
        Z.append(z)
        top.append(hi if isinstance(hi, str) else fmt(hi, 4))
        bottom.append(lo if isinstance(lo, str) else fmt(lo, 4))
        labels.append(f"{names[n]}\n({_goal_word(o)})")
    Z = np.column_stack(Z)
    k = project.k
    s = res["topsis"]
    sn = (s - s.min()) / (np.ptp(s) or 1)
    cmap = plots.cm.viridis
    for r in np.argsort(sn):
        ax.plot(range(Z.shape[1]), Z[r], color=cmap(sn[r]), alpha=0.55, lw=1)
    ax.plot(range(Z.shape[1]), Z[it], color="#ffd400", lw=3.5, label="TOPSIS choice", zorder=5)
    ax.plot(range(Z.shape[1]), Z[it], color="k", lw=0.8, zorder=6)
    if highlight is not None:
        ax.plot(range(Z.shape[1]), Z[highlight], color="#b42318", lw=2.2, label="Selected", zorder=7)
    ax.axvspan(k - 0.5, Z.shape[1] - 0.5, color="#f0f0f0", zorder=0)
    for c in range(Z.shape[1]):
        ax.axvline(c, color="#888", lw=0.6)
        ax.text(c, -0.06, bottom[c], ha="center", va="top", fontsize=7)
        ax.text(c, 1.03, top[c], ha="center", va="bottom", fontsize=7)
    ax.set_xticks(range(Z.shape[1]))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_yticks([])
    ax.set_ylim(-0.15, 1.12)
    ax.set_title("Factors: full design range. Objectives (gray background): top = better.", fontsize=9)
    ax.legend(loc="upper right", fontsize=8)


# ================================================================ ANN vs RSM comparison
def compare_fronts(project, r_rsm, r_ann):
    """Compare the fronts from the RSM and ANN models: front quality (hypervolume), cross-prediction (each front
    re-evaluated with the other model), and the distance between the TOPSIS compromise solutions in factor space."""
    act = r_rsm["active"]
    m = len(act)
    out = {"rsm": r_rsm, "ann": r_ann, "active": act}
    cross = {}
    for key, res in (("rsm", r_rsm), ("ann", r_ann)):
        other = "ann" if key == "rsm" else "rsm"
        Yo = [objective_predict(project, o, res["X"], other) for o in act]
        cross[key] = {"self": res["Yobj"], "other": Yo,
                      "F_self": res["F"], "F_other": np.column_stack([to_F(o, y) for o, y in zip(act, Yo)])}
    out["cross"] = cross
    allF = np.vstack([cross[k][f] for k in cross for f in ("F_self", "F_other")])
    ideal, nadir = allF.min(axis=0), allF.max(axis=0)
    samples = nsga2._hv_samples(m) if m >= 3 else None
    hv = {}
    for k in ("rsm", "ann"):
        for f, ev in (("F_self", k), ("F_other", "ann" if k == "rsm" else "rsm")):
            hv[(k, ev)] = nsga2.hypervolume(cross[k][f], ideal, nadir, samples=samples) / nsga2.hv_max(m)
    out["hv"] = hv
    # prediction difference between the two models, relative to the response data range
    diff = []
    for n, o in enumerate(act):
        y = project.model_data[:, o["index"]]
        vals = np.concatenate([cross[k][s][n] for k in cross for s in ("self", "other")])
        rng_ = float(np.nanmax(y) - np.nanmin(y)) if not o.get("poe") and np.isfinite(y).any() else float(np.ptp(vals))
        rng_ = rng_ if rng_ > 0 else 1.0
        d_all = np.concatenate([np.abs(np.asarray(cross[k]["self"][n]) - np.asarray(cross[k]["other"][n]))
                                for k in cross])
        diff.append({"mean": float(d_all.mean() / rng_ * 100), "max": float(d_all.max() / rng_ * 100),
                     "range": rng_})
    out["diff"] = diff
    lo, hi = project.coded_bounds()
    span = np.where(hi - lo > 0, hi - lo, 1.0)
    xr, xa = r_rsm["X"][r_rsm["i_topsis"]], r_ann["X"][r_ann["i_topsis"]]
    out["dx"] = float(np.max(np.abs(xr - xa) / span) * 100)
    out["dx_factor"] = (np.abs(xr - xa) / span * 100).tolist()
    # model quality for each objective response
    qual = []
    for o in act:
        j = o["index"]
        fit = project.fit(j)
        am = project.ann.get(j)
        pr2 = fit.stats.get("pred_r2", np.nan) if fit is not None else np.nan
        te = am.metrics.get("test", {}).get("r2", np.nan) if am is not None else np.nan
        cv = (am.cv_diag or {}).get("q2", np.nan) if am is not None else np.nan
        qual.append({"j": j, "pred_r2": pr2, "ann_test": te, "ann_q2": cv, "has_ann": am is not None})
    out["quality"] = qual
    return out


def compare_html(project, cmp):
    act = cmp["active"]
    r_rsm, r_ann = cmp["rsm"], cmp["ann"]
    names = [objective_label(project, o) for o in act]
    hv = cmp["hv"]
    h = ["<h2>Pareto Front Comparison: ANN vs RSM</h2>",
         f"<p>The optimization was run twice with the same settings ({ALG_NAME.get(r_rsm['algorithm'])}, population "
         f"{r_rsm['params']['pop']}, seed {r_rsm['params']['seed']}): once using the RSM model for all objectives, "
         "once using ANN.</p>"]
    no_ann = [names[n] for n, q in enumerate(cmp["quality"]) if not q["has_ann"]]
    if no_ann:
        h.append(f"<p class='warn'>No ANN yet for {', '.join(no_ann)}; on the ANN front these responses still use "
                 "RSM.</p>")
    h.append("<h3>Model Quality per Response</h3><table><tr><th class='l'>Objective</th><th>RSM: Pred R²</th>"
             "<th>ANN: test R²</th><th>ANN: cross-validation Q²</th></tr>")
    for n, q in enumerate(cmp["quality"]):
        h.append(f"<tr><td class='l'>{names[n]}</td><td>{fmt(q['pred_r2'], 4)}</td><td>{fmt(q['ann_test'], 4)}</td>"
                 f"<td>{fmt(q['ann_q2'], 4)}</td></tr>")
    h.append("</table>")
    h.append("<h3>Front Size</h3><table><tr><th class='l'></th><th>RSM front</th><th>ANN front</th></tr>"
             f"<tr><td class='l'>Pareto solutions</td><td>{len(r_rsm['X'])}</td><td>{len(r_ann['X'])}</td></tr>"
             f"<tr><td class='l'>Generations run</td><td>{r_rsm['gens_run']}</td><td>{r_ann['gens_run']}</td></tr>"
             f"<tr><td class='l'>Hypervolume, assessed by own model</td><td>{100 * hv[('rsm', 'rsm')]:.1f}%</td>"
             f"<td>{100 * hv[('ann', 'ann')]:.1f}%</td></tr>"
             f"<tr><td class='l'>Hypervolume, assessed by other model</td><td>{100 * hv[('rsm', 'ann')]:.1f}% (by ANN)"
             f"</td><td>{100 * hv[('ann', 'rsm')]:.1f}% (by RSM)</td></tr></table>"
             "<p class='note'>Hypervolume = how much of the objective space the front covers (jointly normalized; "
             "larger is better). 'Assessed by other model': the same front solutions re-predicted with the other "
             "model. If the value drops a lot, the front is only good according to its own model.</p>")
    h.append("<h3>Prediction Difference on Pareto Solutions</h3><table><tr><th class='l'>Objective</th>"
             "<th>Mean difference</th><th>Maximum difference</th></tr>")
    for n, d in enumerate(cmp["diff"]):
        h.append(f"<tr><td class='l'>{names[n]}</td><td>{d['mean']:.1f}%</td><td>{d['max']:.1f}%</td></tr>")
    h.append("</table><p class='note'>As a percentage of the response data range.</p>")
    X = project.to_actual(np.vstack([r_rsm["X"][r_rsm["i_topsis"]], r_ann["X"][r_ann["i_topsis"]]]))
    h.append("<h3>TOPSIS Compromise Solutions</h3><table><tr><th class='l'></th>"
             + "".join(f"<th>{f.name}</th>" for f in project.factors)
             + "".join(f"<th>{nm} (RSM)</th><th>{nm} (ANN)</th>" for nm in names) + "</tr>")
    for row, (lab, key) in enumerate((("RSM front choice", "rsm"), ("ANN front choice", "ann"))):
        res = cmp[key]
        i = res["i_topsis"]
        c = cmp["cross"][key]
        cells = []
        for n in range(len(act)):
            y_r = c["self"][n][i] if key == "rsm" else c["other"][n][i]
            y_a = c["other"][n][i] if key == "rsm" else c["self"][n][i]
            cells.append(f"<td>{fmt(float(y_r), 5)}</td><td>{fmt(float(y_a), 5)}</td>")
        h.append(f"<tr><td class='l'>{lab}</td>"
                 + "".join(f"<td>{project.format_value(f, X[row, f], 5)}</td>" for f in range(project.k))
                 + "".join(cells) + "</tr>")
    h.append(f"</table><p class='note'>Distance between the two choices in factor space: {cmp['dx']:.1f}% of the "
             "design range (factor with the largest difference).</p>")
    # conclusion
    dmean = max(d["mean"] for d in cmp["diff"])
    drop_r = hv[("rsm", "rsm")] - hv[("rsm", "ann")]
    drop_a = hv[("ann", "ann")] - hv[("ann", "rsm")]
    h.append("<h3>Conclusion</h3><ul>")
    # check D: mean difference below 5% for every objective and compromise choices less than 15% of the range apart;
    # name the criterion that failed, so the message matches the rating
    diff_ok, dist_ok = dmean < 5, cmp["dx"] < 15
    diff_txt = (f"the largest mean difference of the predictions on the front is {dmean:.1f}%, "
                + ("within the 5% limit" if diff_ok else "above the 5% limit"))
    dist_txt = (f"the compromise choices are {cmp['dx']:.0f}% of the factor range apart, "
                + ("within the 15% limit" if dist_ok else "above the 15% limit"))
    if diff_ok and dist_ok:
        h.append("<li class='good'>Both models agree: predictions on the front differ by only "
                 f"{dmean:.1f}% on average and the compromise choices are close together. The optimization result "
                 "does not depend on the choice of model.</li>")
    elif dmean < 10:
        if not diff_ok and not dist_ok:
            why = (f"{diff_txt[0].upper() + diff_txt[1:]}, and {dist_txt}. Make confirmation runs at both TOPSIS "
                   "points.")
        elif not diff_ok:
            why = (f"{diff_txt[0].upper() + diff_txt[1:]}, while {dist_txt}. Check the region where the differences "
                   "are largest before relying on the front there.")
        else:
            why = (f"{diff_txt[0].upper() + diff_txt[1:]}, but {dist_txt}. The front is flat in that region: several "
                   "settings give similar results.")
        h.append(f"<li class='warn'>The models agree only partly. {why}</li>")
    else:
        h.append(f"<li class='bad'>The models differ considerably: the largest mean difference of the predictions on "
                 f"the front is {dmean:.1f}%, above the 10% limit for partial agreement, and {dist_txt}. Do not rely "
                 "on just one model: make confirmation runs at both TOPSIS points.</li>")
    better = "ANN" if hv[("ann", "ann")] > hv[("rsm", "rsm")] else "RSM"
    h.append(f"<li>According to each model, the {better} front promises a better trade-off "
             f"(hypervolume {100 * max(hv[('ann', 'ann')], hv[('rsm', 'rsm')]):.1f}% vs "
             f"{100 * min(hv[('ann', 'ann')], hv[('rsm', 'rsm')]):.1f}%).</li>")
    if max(drop_r, drop_a) > 0.05:
        who = "ANN" if drop_a > drop_r else "RSM"
        h.append(f"<li>The {who} front loses a lot of hypervolume when assessed by the other model. Its advantage "
                 f"may exist only in the {who} model (for example, ANN extrapolating or RSM being too simple).</li>")
    q = cmp["quality"]
    pr = [x["pred_r2"] for x in q if np.isfinite(x["pred_r2"])]
    at = [x["ann_test"] for x in q if np.isfinite(x["ann_test"])]
    if pr and at:
        h.append(f"<li>Prediction accuracy on new data: RSM mean Pred R² {np.mean(pr):.3f}, ANN mean test R² "
                 f"{np.mean(at):.3f}. If the fronts differ, the model with higher accuracy is more trustworthy.</li>")
    h.append("</ul>")
    return "".join(h)


def draw_compare(fig, project, cmp, kind, opts=None):
    opts = opts or {}
    act = cmp["active"]
    m = len(act)
    names = [objective_label(project, o) for o in act]
    r_rsm, r_ann = cmp["rsm"], cmp["ann"]
    if kind == "cmp_front":
        a, b = int(opts.get("x", 0)), int(opts.get("y", 1))
        ax = fig.add_subplot(111)
        for key, res, col, mk, lab in (("rsm", r_rsm, C_RSM, "o", "RSM front"), ("ann", r_ann, C_ANN, "^", "ANN front")):
            Y = np.column_stack(res["Yobj"])
            F = res["F"]
            pf = pair_front(F[:, a], F[:, b]) if a != b else np.array([], int)
            if len(pf) > 1:
                ax.plot(Y[pf, a], Y[pf, b], "-", color=col, lw=1, alpha=0.6)
            ax.scatter(Y[:, a], Y[:, b], marker=mk, s=34, color=col, alpha=0.75, edgecolor="white", lw=0.4,
                       label=f"{lab} ({len(Y)} solutions)", zorder=3)
            i = res["i_topsis"]
            c = cmp["cross"][key]
            ya, yb = c["other"][a][i], c["other"][b][i]
            ax.scatter(Y[i, a], Y[i, b], marker="*", s=360, color=col, edgecolor="k", zorder=6,
                       label=f"TOPSIS {lab.split()[0]}")
            ax.scatter(ya, yb, marker="*", s=300, facecolor="white", edgecolor=col, lw=1.5, zorder=5)
            ax.annotate("", (ya, yb), (Y[i, a], Y[i, b]),
                        arrowprops={"arrowstyle": "->", "color": col, "lw": 1, "ls": "--"}, zorder=4)
        ax.scatter([], [], marker="*", s=200, facecolor="white", edgecolor="#555", lw=1.2,
                   label="TOPSIS choice predicted\nby the other model")
        ax.set_xlabel(f"{names[a]} ({_goal_word(act[a])})")
        ax.set_ylabel(f"{names[b]} ({_goal_word(act[b])})")
        ax.set_title("RSM front vs ANN front" + (f" (2D projection of {m} objectives)" if m > 2 else "")
                     + "\nHollow stars: TOPSIS choice when predicted with the other model", fontsize=9)
        ax.grid(alpha=0.3)
        ax.legend(fontsize=8, loc="best")
    elif kind == "cmp_cross":
        nc = min(3, m)
        nr = int(np.ceil(m / nc))
        axs = np.atleast_1d(fig.subplots(nr, nc, squeeze=False)).ravel()
        for n in range(m):
            ax = axs[n]
            for key, col, mk, lab in (("rsm", C_RSM, "o", "RSM front solutions"), ("ann", C_ANN, "^", "ANN front solutions")):
                c = cmp["cross"][key]
                yr = c["self"][n] if key == "rsm" else c["other"][n]
                ya = c["other"][n] if key == "rsm" else c["self"][n]
                ax.scatter(yr, ya, s=16, marker=mk, color=col, alpha=0.7, label=lab)
            lo = min(ax.get_xlim()[0], ax.get_ylim()[0])
            hi = max(ax.get_xlim()[1], ax.get_ylim()[1])
            ax.plot([lo, hi], [lo, hi], color="#888", lw=1, ls="--")
            ax.set_xlim(lo, hi)
            ax.set_ylim(lo, hi)
            ax.set_title(f"{names[n]}: difference {cmp['diff'][n]['mean']:.1f}%", fontsize=8)
            ax.set_xlabel("RSM prediction", fontsize=8)
            ax.set_ylabel("ANN prediction", fontsize=8)
            ax.tick_params(labelsize=7)
            ax.grid(alpha=0.3)
            if n == 0:
                ax.legend(fontsize=7)
        for ax in axs[m:]:
            ax.axis("off")
        fig.suptitle("Cross-prediction: each Pareto solution predicted with both models (dashed line = agreement)",
                     fontsize=10)
    else:  # cmp_factors
        ax = fig.add_subplot(111)
        (Zr, Za), top, bottom, labels = _factor_axes(project, [project.to_actual(r_rsm["X"]),
                                                               project.to_actual(r_ann["X"])])
        Zr, Za = np.column_stack(Zr), np.column_stack(Za)
        xs = range(Zr.shape[1])
        for Z, col in ((Zr, C_RSM), (Za, C_ANN)):
            for r in range(len(Z)):
                ax.plot(xs, Z[r], color=col, alpha=0.25, lw=0.8)
        ax.plot(xs, Zr[r_rsm["i_topsis"]], color=C_RSM, lw=3, label="TOPSIS RSM front")
        ax.plot(xs, Za[r_ann["i_topsis"]], color=C_ANN, lw=3, label="TOPSIS ANN front")
        for c in xs:
            ax.axvline(c, color="#888", lw=0.6)
            ax.text(c, -0.06, bottom[c], ha="center", va="top", fontsize=7)
            ax.text(c, 1.03, top[c], ha="center", va="bottom", fontsize=7)
        ax.set_xticks(list(xs))
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_yticks([])
        ax.set_ylim(-0.15, 1.12)
        ax.set_title("Factor settings of Pareto solutions: blue = RSM front, orange = ANN front (full design range)",
                     fontsize=9)
        ax.legend(loc="upper right", fontsize=8)


# ================================================================ convergence & adequacy report
def convergence_html(project, res, adequacy=None):
    chk = res.get("check") or {}
    alg = ALG_NAME.get(res.get("algorithm"), "NSGA-II")
    prm = res["params"]
    h = [f"<h2>Convergence & Adequacy - {alg}</h2>",
         f"<p>Population {prm['pop']}, maximum generations {prm['gens']}, {res.get('gens_run')} generations run"
         + (f", {res['stop_reason']}" if res.get("stop_reason") else "") + ".</p>"]
    if chk.get("converged_at"):
        h.append(f"<p>The convergence criterion was first met at generation <b>{chk['converged_at']}</b>.</p>")
    h.append("<h3>Quick Assessment</h3><ul>")
    for n in chk.get("notes", []):
        cls = "good" if n.startswith("Enough generations") else ("warn" if n.startswith("Not enough generations") else "")
        h.append(f"<li class='{cls}'>{n}</li>")
    h.append("</ul>")
    h.append("<h3>How to Read</h3><ul>"
             "<li><b>Enough generations</b> when the hypervolume curve has flattened: in the last window of "
             "generations, the mean hypervolume (a measure of front quality combining closeness and spread) hardly "
             "increases compared with the previous window, and the ideal point (best value of each objective) does "
             "not shift. Automatic stop uses this criterion so the run does not go on longer than needed.</li>"
             "<li><b>Enough population</b> when repeating the optimization with another seed gives the same "
             "hypervolume (difference below 1%) and doubling the population no longer increases the hypervolume. "
             "Click <b>Check Population Size</b> to test this.</li></ul>")
    if adequacy:
        v = adequacy["verdict"]
        h.append("<h3>Adequacy Check Results</h3><table><tr><th class='l'>Run</th><th>Population</th><th>Seed</th>"
                 "<th>Generations</th><th>Pareto solutions</th><th>Hypervolume</th><th>Difference vs current</th>"
                 "<th>TOPSIS shift (% of factor range)</th></tr>")
        hv0 = adequacy["runs"][0]["hv"]
        for r in adequacy["runs"]:
            d = (r["hv"] - hv0) / max(hv0, 1e-12) * 100
            h.append(f"<tr><td class='l'>{r['label']}</td><td>{r['pop']}</td><td>{r['seed']}</td>"
                     f"<td>{r['gens']}{' (converged)' if r['stop'] else ''}</td><td>{r['n']}</td>"
                     f"<td>{100 * r['hv']:.2f}%</td><td>{d:+.2f}%</td><td>{r['dx'] * 100:.1f}</td></tr>")
        h.append("</table><ul>")
        h.append(f"<li class='{'good' if v['gens_ok'] else 'warn'}'>Generations: "
                 + ("sufficient (the hypervolume has flattened)." if v["gens_ok"] else
                    "not sufficient, increase the maximum generations.") + "</li>")
        if v["pop_ok"]:
            h.append(f"<li class='good'>Population: sufficient. Another seed changes the hypervolume by only "
                     f"{100 * v['spread']:.2f}% and a 2x population adds only {100 * v['gain_big']:.2f}%.</li>")
        else:
            why = []
            if v["spread"] >= 0.01:
                why.append(f"the result changes by {100 * v['spread']:.1f}% when the seed is changed")
            if v["gain_big"] >= 0.01:
                why.append(f"a 2x population increases the hypervolume by {100 * v['gain_big']:.1f}%")
            h.append(f"<li class='warn'>Population: not sufficient ({'; '.join(why)}). Increase the population, "
                     f"for example to {2 * prm['pop']}, then check again.</li>")
        if v["dx_max"] > 0.15:
            h.append(f"<li>The TOPSIS choice shifts by up to {100 * v['dx_max']:.0f}% of the factor range between "
                     "runs. The front tends to be flat around the compromise: many settings give nearly the same "
                     "result.</li>")
        h.append("</ul>")
    return "".join(h)


# ================================================================ tab
class NsgaTab(QWidget):
    send_to_prediction = Signal(object)
    model_changed = Signal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.project = None
        self.result = None
        self.compare = None
        self.adequacy = None
        self.task = None
        self._rows = []
        self._syncing = False

        lay = QVBoxLayout(self)
        g = QGroupBox("Objectives && Constraints (Select at Least 2 Objectives)")
        gl = QVBoxLayout(g)
        self.tbl = QTableWidget(0, len(COLS))
        self.tbl.setHorizontalHeaderLabels(COLS)
        self.tbl.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.tbl.verticalHeader().setVisible(False)
        gl.addWidget(self.tbl)

        prm = QHBoxLayout()
        self.sp_pop = QSpinBox()
        self.sp_pop.setRange(20, 2000)
        self.sp_pop.setValue(100)
        self.sp_gen = QSpinBox()
        self.sp_gen.setRange(10, 10000)
        self.sp_gen.setValue(300)
        self.sp_pc = QDoubleSpinBox()
        self.sp_pc.setRange(0.1, 1.0)
        self.sp_pc.setSingleStep(0.05)
        self.sp_pc.setValue(0.9)
        self.sp_etac = QSpinBox()
        self.sp_etac.setRange(1, 100)
        self.sp_etac.setValue(15)
        self.sp_etam = QSpinBox()
        self.sp_etam.setRange(1, 100)
        self.sp_etam.setValue(20)
        self.sp_seed = QSpinBox()
        self.sp_seed.setRange(0, 999999)
        self.sp_seed.setValue(1)
        for lab, w in (("Population:", self.sp_pop), ("Max generations:", self.sp_gen), ("P crossover:", self.sp_pc),
                       ("η crossover:", self.sp_etac), ("η mutation:", self.sp_etam), ("Seed:", self.sp_seed)):
            prm.addWidget(QLabel(lab))
            prm.addWidget(w)
        prm.addStretch()
        gl.addLayout(prm)

        prm2 = QHBoxLayout()
        self.cb_alg = QComboBox()
        for key, label in nsga2.ALGORITHMS.items():
            self.cb_alg.addItem(label, key)
        self.cb_alg.setToolTip("NSGA-II uses crowding distance; good for 2 objectives.\n"
                               "NSGA-III uses reference points so solutions stay evenly spread for 3 or more\n"
                               "objectives (when almost all solutions are mutually non-dominated).")
        self.chk_stop = QCheckBox("Stop automatically on convergence")
        self.chk_stop.setChecked(True)
        self.chk_stop.setToolTip("Stop before the maximum generations when the front hypervolume hardly increases\n"
                                 "and the ends of the front do not shift during the last window of generations.")
        self.sp_tol = QDoubleSpinBox()
        self.sp_tol.setRange(0.001, 5.0)
        self.sp_tol.setDecimals(3)
        self.sp_tol.setSingleStep(0.05)
        self.sp_tol.setValue(0.1)
        self.sp_tol.setSuffix(" %")
        self.sp_tol.setToolTip("Maximum hypervolume change within the generation window to count as converged.")
        self.sp_win = QSpinBox()
        self.sp_win.setRange(5, 500)
        self.sp_win.setValue(20)
        self.sp_win.setSuffix(" generations")
        prm2.addWidget(QLabel("Algorithm:"))
        prm2.addWidget(self.cb_alg)
        prm2.addSpacing(12)
        prm2.addWidget(self.chk_stop)
        prm2.addWidget(QLabel("tolerance:"))
        prm2.addWidget(self.sp_tol)
        prm2.addWidget(QLabel("within"))
        prm2.addWidget(self.sp_win)
        prm2.addStretch()
        gl.addLayout(prm2)

        btns = QHBoxLayout()
        self.btn_run = QPushButton("Run Optimization")
        self.btn_cmp = QPushButton("Compare ANN vs RSM")
        self.btn_cmp.setToolTip("Run the optimization twice with the same settings: all objectives using RSM, then\n"
                                "all using ANN. The fronts, cross-predictions and compromise solutions are compared.")
        self.btn_adq = QPushButton("Check Population Size")
        self.btn_adq.setToolTip("Repeat the optimization with 2 other seeds and with a 2x population. If the\n"
                                "hypervolume does not change (< 1%), the population and generations are sufficient.")
        self.btn_stop = QPushButton("Stop")
        self.btn_stop.setVisible(False)
        self.progress = QProgressBar()
        self.progress.setMaximumWidth(220)
        self.lbl_prog = QLabel()
        self.lbl_prog.setStyleSheet("color: #555;")
        for b in (self.btn_run, self.btn_cmp, self.btn_adq):
            b.setMinimumWidth(b.fontMetrics().horizontalAdvance(b.text()) + 40)
            btns.addWidget(b)
        btns.addWidget(self.btn_stop)
        btns.addWidget(self.progress)
        btns.addWidget(self.lbl_prog, 1)
        gl.addLayout(btns)
        note = QLabel("Factor limits are taken from the Optimization tab. Response lower/upper limits (optional) "
                      "become constraints. TOPSIS weights are used to select the compromise solution. Choose the model "
                      "for each response (RSM or ANN) directly in the Model column.")
        note.setWordWrap(True)
        note.setStyleSheet("color: #555;")
        gl.addWidget(note)
        lay.addWidget(g, 2)

        split = QSplitter(Qt.Horizontal)
        left = QWidget()
        lv = QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        vrow = QHBoxLayout()
        vrow.addWidget(QLabel("Show:"))
        self.cb_view = QComboBox()
        self.cb_view.addItem("Optimization results", "main")
        vrow.addWidget(self.cb_view, 1)
        lv.addLayout(vrow)
        self.lbl_best = QLabel()
        self.lbl_best.setWordWrap(True)
        self.lbl_best.setStyleSheet("font-weight: bold; color: #1f5f99;")
        lv.addWidget(self.lbl_best)
        self.tbl_sol = QTableWidget()
        install_table_copy(self.tbl_sol)
        self.tbl_sol.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tbl_sol.setSelectionBehavior(QTableWidget.SelectRows)
        self.tbl_sol.verticalHeader().setVisible(False)
        self.tbl_sol.setSortingEnabled(False)
        lv.addWidget(self.tbl_sol)
        bb = QHBoxLayout()
        self.btn_send = QPushButton("Send Selected Solution to Prediction")
        self.btn_export = QPushButton("Export Pareto to Excel...")
        bb.addWidget(self.btn_send)
        bb.addWidget(self.btn_export)
        lv.addLayout(bb)
        split.addWidget(left)

        self.rtabs = QTabWidget()
        right = QWidget()
        rv = QVBoxLayout(right)
        rv.setContentsMargins(4, 4, 4, 4)
        pf = QFormLayout()
        self.cb_plot = QComboBox()
        pf.addRow("Graph:", self.cb_plot)
        axes = QHBoxLayout()
        self.cb_x, self.cb_y, self.cb_c, self.cb_s = QComboBox(), QComboBox(), QComboBox(), QComboBox()
        self.axis_labels = []
        for lab, cb in (("X:", self.cb_x), ("Y:", self.cb_y), ("Color:", self.cb_c), ("Size:", self.cb_s)):
            lb = QLabel(lab)
            self.axis_labels.append(lb)
            axes.addWidget(lb)
            axes.addWidget(cb, 1)
        self.axes_row = QWidget()
        self.axes_row.setLayout(axes)
        axes.setContentsMargins(0, 0, 0, 0)
        pf.addRow(self.axes_row)
        rv.addLayout(pf)
        self.canvas = MplCanvas(figsize=(6, 5))
        rv.addWidget(self.canvas)
        self.rtabs.addTab(right, "Graph")
        self.txt_conv = ReportBrowser()
        self.rtabs.addTab(self.txt_conv, "Convergence && Adequacy")
        self.txt_cmp = ReportBrowser()
        self.rtabs.addTab(self.txt_cmp, "ANN vs RSM")
        split.addWidget(self.rtabs)
        split.setSizes([620, 600])
        lay.addWidget(split, 3)

        self.btn_run.clicked.connect(self.run)
        self.btn_cmp.clicked.connect(self.run_compare)
        self.btn_adq.clicked.connect(self.run_adequacy)
        self.btn_stop.clicked.connect(self.stop)
        self.btn_send.clicked.connect(self.send)
        self.btn_export.clicked.connect(self.export)
        self.cb_plot.currentIndexChanged.connect(self._on_plot_kind)
        for cb in (self.cb_x, self.cb_y, self.cb_c, self.cb_s):
            cb.currentIndexChanged.connect(self.draw)
        self.cb_view.currentIndexChanged.connect(self._on_view)
        self.tbl_sol.itemSelectionChanged.connect(self.draw)

    # ---------------------------------------------------------------
    def set_project(self, project):
        if self.task is not None:
            self.task.cancel()
        self.project = project
        self.result = project.nsga_result
        self.compare = getattr(project, "nsga_compare", None)
        self.adequacy = None
        self.fill_table()
        self._fill_views()
        self.show_result()

    def refresh(self):
        if self.project is not None:
            self._refresh_sources()

    def _refresh_sources(self):
        """Model column: RSM/ANN choice per response; POE rows follow the model of their response."""
        p = self.project
        for n, (j, poe) in enumerate(self._rows):
            if poe:
                it = self.tbl.item(n, 1)
                if it:
                    it.setText(("ANN" if p.source(j) == "ann" else "RSM") + " + POE")
            else:
                self.tbl.setCellWidget(n, 1, model_source_combo(p, j, self._on_source))

    def _on_source(self):
        self._refresh_pos()
        self.model_changed.emit()

    def _refresh_pos(self):
        p = self.project
        for n, (j, poe) in enumerate(self._rows):
            if poe and self.tbl.item(n, 1):
                self.tbl.item(n, 1).setText(("ANN" if p.source(j) == "ann" else "RSM") + " + POE")

    def fill_table(self):
        p = self.project
        self._rows = [(j, False) for j in range(len(p.responses))]
        if p.has_poe:
            self._rows += [(j, True) for j in range(len(p.responses)) if p.analysis_kind(j) != "glm"]
        self.tbl.setRowCount(len(self._rows))
        for n, (j, poe) in enumerate(self._rows):
            name = QTableWidgetItem(("POE " if poe else "") + p.response_label(j))
            name.setFlags(name.flags() & ~Qt.ItemIsEditable)
            self.tbl.setItem(n, 0, name)
            src = QTableWidgetItem()
            src.setFlags(src.flags() & ~Qt.ItemIsEditable)
            self.tbl.setItem(n, 1, src)
            cb = QComboBox()
            for key, label in GOALS:
                cb.addItem(label, key)
            prev = p.opt_criteria.get(j, {}).get("goal", "none") if not poe else "none"
            cb.setCurrentIndex(max(cb.findData(prev if prev in ("maximize", "minimize", "target") else "none"), 0))
            self.tbl.setCellWidget(n, 2, cb)
            tgt = p.opt_criteria.get(j, {}).get("target", "")
            for c, v in ((3, "" if poe or tgt == "" else fmt(tgt, 8)), (4, ""), (5, ""), (6, "1")):
                self.tbl.setItem(n, c, QTableWidgetItem(v))
        self._refresh_sources()

    def _objectives(self):
        objs = []
        for n, (j, poe) in enumerate(self._rows):
            goal = self.tbl.cellWidget(n, 2).currentData()

            def num(c, default=np.nan):
                t = self.tbl.item(n, c).text().strip() if self.tbl.item(n, c) else ""
                return parse_float(t) if t else default

            o = {"index": j, "goal": goal, "poe": poe, "target": num(3), "low": num(4), "high": num(5),
                 "weight": num(6, 1.0)}
            if goal == "target" and not np.isfinite(o["target"]):
                raise ValueError(f"Enter a target value for {self.tbl.item(n, 0).text()}.")
            if goal != "none" or np.isfinite(o["low"]) or np.isfinite(o["high"]):
                objs.append(o)
        return objs

    def _settings(self):
        p = self.project
        objs = self._objectives()
        if sum(o["goal"] != "none" for o in objs) < 2:
            raise ValueError("Select at least 2 objectives (maximize/minimize/target).")
        bounds = {i: (c["low"], c["high"]) for key, c in p.opt_criteria.items()
                  if isinstance(key, str) and key.startswith("f") and key[1:].isdigit()
                  for i in [int(key[1:])] if i < p.k and not p.factors[i].categoric}
        kw = {"pop": self.sp_pop.value(), "gens": self.sp_gen.value(), "seed": self.sp_seed.value(),
              "pc": self.sp_pc.value(), "eta_c": self.sp_etac.value(), "eta_m": self.sp_etam.value(),
              "algorithm": self.cb_alg.currentData(), "auto_stop": self.chk_stop.isChecked(),
              "stop_tol": self.sp_tol.value() / 100, "stop_window": self.sp_win.value()}
        return objs, bounds, kw

    # ---------------------------------------------------------------- background work
    def _set_busy(self, busy):
        for w in (self.btn_run, self.btn_cmp, self.btn_adq, self.tbl, self.btn_send, self.btn_export):
            w.setEnabled(not busy)
        self.btn_stop.setVisible(busy)
        if not busy:
            self.lbl_prog.setText("")

    def _start(self, job, on_done):
        proj = self.project
        self.task = Task(job, self)
        self.task.progressed.connect(self._on_progress)
        self.task.done.connect(lambda res: self._finish(res, on_done, proj))
        self.task.failed.connect(self._on_failed)
        self._set_busy(True)
        self.progress.setValue(0)
        self.task.start()

    def _on_progress(self, text, i, n):
        self.lbl_prog.setText(text)
        if n:
            self.progress.setValue(int(100 * i / n))

    def _finish(self, res, on_done, proj):
        self._set_busy(False)
        self.task = None
        self.progress.setValue(100)
        if self.project is proj:
            on_done(res)

    def _on_failed(self, msg):
        self._set_busy(False)
        self.task = None
        QMessageBox.warning(self, "Multi-Objective Optimization", msg)

    def stop(self):
        if self.task is not None:
            self.lbl_prog.setText("Stopping...")
            self.task.cancel()

    def _runner(self, objs, bounds, kw, source=None, label=""):
        p = self.project

        def one(report, cancelled, stage=""):
            def prog(g, total):
                if g % 5 == 0 or g == total:
                    report(f"{stage}{label}generation {g} of max {total}", g, total)
            return p.run_nsga(objs, bounds, progress=prog, source=source, store=False, cancel=cancelled, **kw)
        return one

    def run(self):
        try:
            objs, bounds, kw = self._settings()
        except ValueError as exc:
            QMessageBox.warning(self, "Multi-Objective Optimization", str(exc))
            return
        one = self._runner(objs, bounds, kw)

        def done(res):
            self.project.nsga_result = res
            self.result = res
            self.adequacy = None
            self.cb_view.setCurrentIndex(0)
            self.show_result()

        self._start(lambda report, cancelled: one(report, cancelled), done)

    def run_compare(self):
        p = self.project
        try:
            objs, bounds, kw = self._settings()
        except ValueError as exc:
            QMessageBox.warning(self, "Multi-Objective Optimization", str(exc))
            return
        act = [o for o in objs if o["goal"] != "none"]
        if not any(o["index"] in p.ann for o in act):
            QMessageBox.information(self, "ANN vs RSM", "There is no ANN model for the objective responses yet. Train "
                                    "an ANN first on the ANN page (Optimize All Responses button).")
            return
        run_r = self._runner(objs, bounds, kw, "rsm")
        run_a = self._runner(objs, bounds, kw, "ann")

        def job(report, cancelled):
            r = run_r(report, cancelled, "[1/2] RSM: ")
            a = run_a(report, cancelled, "[2/2] ANN: ")
            return compare_fronts(p, r, a)

        def done(cmp):
            p.nsga_compare = cmp
            self.compare = cmp
            self._fill_views()
            self.txt_cmp.setHtml(REPORT_CSS + compare_html(p, cmp))
            self._fill_plot_kinds("cmp_front")
            self.rtabs.setCurrentIndex(0)
            if self.result is None:
                self.show_result()
            else:
                self.draw()

        self._start(job, done)

    def run_adequacy(self):
        res = self.view_result()
        if not res:
            QMessageBox.information(self, "Population Size Check", "Run the optimization first.")
            return
        p = self.project

        def job(report, cancelled):
            return p.nsga_adequacy(res, progress=lambda k, n, lab: report(f"Population check {k + 1}/{n}: {lab}", k, n),
                                   cancel=cancelled)

        def done(adq):
            self.adequacy = adq
            self.show_convergence()
            self.rtabs.setCurrentWidget(self.txt_conv)

        self._start(job, done)

    # ---------------------------------------------------------------- display
    def _fill_views(self):
        self._syncing = True
        cur = self.cb_view.currentData()
        self.cb_view.clear()
        self.cb_view.addItem("Optimization results", "main")
        if self.compare:
            self.cb_view.addItem("Comparison: RSM front", "rsm")
            self.cb_view.addItem("Comparison: ANN front", "ann")
        i = self.cb_view.findData(cur)
        self.cb_view.setCurrentIndex(max(i, 0))
        self._syncing = False

    def view_result(self):
        key = self.cb_view.currentData()
        if key in ("rsm", "ann") and self.compare:
            return self.compare[key]
        return self.result if self.result else (self.compare["ann"] if self.compare else None)

    def _on_view(self):
        if not self._syncing:
            self.show_result()

    def _fill_plot_kinds(self, prefer=None):
        res = self.view_result()
        self._syncing = True
        cur = prefer or self.cb_plot.currentData()
        self.cb_plot.clear()
        for key, label in KINDS + (CMP_KINDS if self.compare else []):
            self.cb_plot.addItem(label, key)
        if cur is None and res:
            cur = "front" if len(res["active"]) == 2 else "front2d"
        i = self.cb_plot.findData(cur)
        self.cb_plot.setCurrentIndex(max(i, 0))
        # axes
        act = (res or {}).get("active") or []
        names = [objective_label(self.project, o) for o in act]
        for cb in (self.cb_x, self.cb_y, self.cb_c, self.cb_s):
            prev = cb.currentData()
            cb.clear()
            if cb is self.cb_c:
                cb.addItem("TOPSIS score", "topsis")
            if cb is self.cb_s:
                cb.addItem("(uniform)", None)
            for k, nm in enumerate(names):
                cb.addItem(nm, k)
            j = cb.findData(prev)
            if j >= 0 and prev is not None:
                cb.setCurrentIndex(j)
        if len(names) >= 2 and self.cb_x.currentData() == self.cb_y.currentData():
            self.cb_x.setCurrentIndex(0)
            self.cb_y.setCurrentIndex(1)
            if len(names) >= 3:
                self.cb_c.setCurrentIndex(self.cb_c.findData(2))
            if len(names) >= 4:
                self.cb_s.setCurrentIndex(self.cb_s.findData(3))
        self._syncing = False
        self._axes_visibility()

    def _axes_visibility(self):
        kind = self.cb_plot.currentData()
        show_xy = kind in ("front2d", "cmp_front")
        self.axes_row.setVisible(show_xy)
        for w in (self.cb_c, self.cb_s, self.axis_labels[2], self.axis_labels[3]):
            w.setVisible(kind == "front2d")

    def _on_plot_kind(self):
        if self._syncing:
            return
        self._axes_visibility()
        self.draw()

    def show_result(self):
        p = self.project
        res = self.view_result()
        self.tbl_sol.clear()
        self._fill_plot_kinds()
        if not res:
            self.tbl_sol.setRowCount(0)
            self.tbl_sol.setColumnCount(0)
            self.lbl_best.setText("")
            self.canvas.message("Set the objectives, then click Run Optimization.")
            self.txt_conv.setHtml(REPORT_CSS + "<h2>Convergence & Adequacy</h2><p>No results yet.</p>")
            self.txt_cmp.setHtml(REPORT_CSS + "<h2>ANN vs RSM</h2><p>Click <b>Compare ANN vs RSM</b> to "
                                 "run the optimization with both models and compare their Pareto fronts.</p>")
            return
        X = p.to_actual(res["X"])
        heads = ["No", "Mark"] + [f.name for f in p.factors] + [r.name for r in p.responses] + ["TOPSIS score"]
        self.tbl_sol.setColumnCount(len(heads))
        self.tbl_sol.setHorizontalHeaderLabels(heads)
        self.tbl_sol.setRowCount(len(X))
        it, ik = res["i_topsis"], res["i_knee"]
        for r in range(len(X)):
            mark = ("★ TOPSIS " if r == it else "") + ("◆ Knee" if r == ik else "")
            vals = [str(r + 1), mark.strip()] + [p.format_value(i, X[r, i], 5) for i in range(p.k)] + \
                [fmt(float(res["Y"][j][r]), 5) if j in res["Y"] else "-" for j in range(len(p.responses))] + \
                [f"{res['topsis'][r]:.4f}"]
            for c, v in enumerate(vals):
                cell = QTableWidgetItem(v)
                cell.setTextAlignment(Qt.AlignCenter)
                if r == it:
                    cell.setBackground(QBrush(QColor("#fff3bf")))
                elif r == ik:
                    cell.setBackground(QBrush(QColor("#ffe8d9")))
                self.tbl_sol.setItem(r, c, cell)
        self.tbl_sol.resizeColumnsToContents()
        desc = ", ".join(f"{p.factors[i].name} = {p.format_value(i, X[it, i], 5)}" for i in range(p.k))
        resp = ", ".join(f"{p.responses[j].name} = {fmt(float(res['Y'][j][it]), 5)}" for j in res["Y"])
        alg = ALG_NAME.get(res.get("algorithm"), "NSGA-II")
        stop = f", automatic stop at generation {res['stop_gen']}" if res.get("stop_gen") and \
            "converged" in (res.get("stop_reason") or "") else f", {res.get('gens_run', '')} generations"
        self.lbl_best.setText(f"{alg}: {len(X)} Pareto solutions{stop}. Best compromise (TOPSIS): {desc} → {resp}")
        self.show_convergence()
        if self.compare:
            self.txt_cmp.setHtml(REPORT_CSS + compare_html(p, self.compare))
        else:
            self.txt_cmp.setHtml(REPORT_CSS + "<h2>ANN vs RSM</h2><p>Click <b>Compare ANN vs RSM</b> to "
                                 "run the optimization with both models and compare their Pareto fronts.</p>")
        self.draw()

    def show_convergence(self):
        res = self.view_result()
        if res:
            adq = self.adequacy if self.adequacy and self.cb_view.currentData() == "main" else None
            self.txt_conv.setHtml(REPORT_CSS + convergence_html(self.project, res, adq))

    def selected(self):
        rows = self.tbl_sol.selectionModel().selectedRows() if self.tbl_sol.selectionModel() else []
        return rows[0].row() if rows else None

    def plot_opts(self):
        return {"x": self.cb_x.currentData() or 0, "y": self.cb_y.currentData() if self.cb_y.currentData()
                is not None else 1, "color": self.cb_c.currentData(), "size": self.cb_s.currentData()}

    def draw(self):
        if self._syncing:
            return
        res = self.view_result()
        if not res:
            return
        fig = self.canvas.figure
        fig.clear()
        kind = self.cb_plot.currentData() or "front"
        if kind.startswith("cmp_") and self.compare:
            draw_compare(fig, self.project, self.compare, kind, self.plot_opts())
        else:
            draw_pareto(fig, self.project, res, self.selected(), kind, self.plot_opts())
        self.canvas.draw()

    def send(self):
        res = self.view_result()
        if not res:
            QMessageBox.information(self, "Multi-Objective Optimization", "Run the optimization first.")
            return
        r = self.selected()
        self.send_to_prediction.emit(res["X"][res["i_topsis"] if r is None else r])

    def export(self):
        p, res = self.project, self.view_result()
        if not res:
            return
        path, _ = QFileDialog.getSaveFileName(self, "Export Pareto", "pareto.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        write_pareto_excel(p, res, path, self.compare)


def write_pareto_excel(p, res, path, compare=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def head(ws):
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EEF6")

    def front_sheet(ws, r_):
        X = p.to_actual(r_["X"])
        ws.append(["No", "Mark"] + [p.factor_label(i) for i in range(p.k)]
                  + [p.response_label(j) for j in range(len(p.responses))] + ["TOPSIS score"])
        head(ws)
        for r in range(len(X)):
            mark = ("TOPSIS " if r == r_["i_topsis"] else "") + ("Knee" if r == r_["i_knee"] else "")
            ws.append([r + 1, mark.strip()]
                      + [p.format_value(i, X[r, i]) if p.factors[i].categoric else float(X[r, i]) for i in range(p.k)]
                      + [float(r_["Y"][j][r]) if j in r_["Y"] else None for j in range(len(p.responses))]
                      + [float(r_["topsis"][r])])

    wb = Workbook()
    ws = wb.active
    ws.title = "Pareto"
    front_sheet(ws, res)
    ws2 = wb.create_sheet("Settings")
    for k, v in res["params"].items():
        ws2.append([k, v])
    ws2.append(["generations run", res.get("gens_run")])
    ws2.append(["stop reason", res.get("stop_reason") or "maximum generations"])
    ws2.append([])
    ws2.append(["Objective", "Goal", "Target", "Lower limit", "Upper limit", "Weight"])
    for o in res["objectives"]:
        ws2.append([objective_label(p, o), o["goal"]] + [None if not np.isfinite(o.get(k, np.nan)) else o[k]
                                                         for k in ("target", "low", "high", "weight")])
    ws3 = wb.create_sheet("Convergence")
    ws3.append(["Generation", "Solutions in front", "Feasible solutions", "Hypervolume change (%)",
                "Front end shift (%)"])
    head(ws3)
    for h in res["history"]:
        g, s = h.get("gain", np.nan), h.get("shift", np.nan)
        ws3.append([h["gen"], h["n_front"], h["n_feasible"], None if not np.isfinite(g) else 100 * g,
                    None if not np.isfinite(s) else 100 * s])
    if res.get("hv_curve"):
        ws3.append([])
        ws3.append(["Generation", "Hypervolume (% of max)"])
        for g, v in res["hv_curve"]:
            ws3.append([g, 100 * v])
    if compare:
        front_sheet(wb.create_sheet("RSM Front"), compare["rsm"])
        front_sheet(wb.create_sheet("ANN Front"), compare["ann"])
        ws4 = wb.create_sheet("ANN vs RSM")
        names = [objective_label(p, o) for o in compare["active"]]
        ws4.append(["Objective", "Mean difference (% of range)", "Max difference (% of range)"])
        head(ws4)
        for n, d in enumerate(compare["diff"]):
            ws4.append([names[n], d["mean"], d["max"]])
        ws4.append([])
        ws4.append(["Front", "Assessed by model", "Hypervolume (% of max)"])
        for (k, ev), v in compare["hv"].items():
            ws4.append([k.upper(), ev.upper(), 100 * v])
    wb.save(path)
