"""Report builder: HTML (screen & PDF), PDF export, Excel export."""
import datetime
import io

import numpy as np
from scipy import stats
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

from doe import designs, models

from . import plots
from .common import REPORT_CSS, fmt, fmt_p, p_class

STAT_LABELS = [
    ("std_dev", "Std. Dev."), ("mean", "Mean"), ("cv", "C.V. %"), ("r2", "R²"),
    ("adj_r2", "Adjusted R²"), ("pred_r2", "Predicted R²"), ("press", "PRESS"),
    ("adeq_precision", "Adeq Precision"),
]


def response_lhs(project, j):
    r = project.responses[j]
    fam = project.response_family(j)
    if fam == "binomial":
        return f"logit(p {r.name})"
    if fam == "poisson":
        return f"ln(λ {r.name})"
    tr = project.model_spec(j).get("transform") or {"kind": "none"}
    name = r.name
    k = float(tr.get("shift", 0) or 0)
    if k:
        name = f"({name} {k:+g})"
    return {"none": name, "sqrt": f"Sqrt({name})", "ln": f"Ln({name})", "log10": f"Log10({name})",
            "inverse": f"1/({name})",
            "power": f"({name})^{tr.get('lam', 1):g}"}.get(tr.get("kind", "none"), name)


# ------------------------------------------------------------------ analysis HTML
def summary_html(project, j, rows, suggested, heading="h2"):
    h = [f"<{heading}>Fit Summary - {project.response_label(j)}</{heading}>"]
    if not rows:
        return "".join(h) + ("<p class='note'>Fit Summary is not available for split-plot (REML) analysis. "
                             "Choose a model, then use the Wald tests on the ANOVA and Backward Elimination tabs.</p>")
    glm = project.analysis_kind(j) == "glm"
    if glm:
        h.append("<table><tr><th class='l'>Source</th><th>Sequential p (LR)</th><th>Deviance</th>"
                 "<th>AIC</th><th>McFadden R²</th><th class='l'>Remarks</th></tr>")
        for r in rows[1:]:
            f = r["fit"]
            note = []
            if r["order"] == suggested:
                note.append("<span class='sig'>Suggested</span>")
            if r["aliased"]:
                note.append("<span class='ns'>Aliased</span>")
            h.append(f"<tr><td class='l'>{models.order_label(r['order'])}</td>"
                     f"<td class='{p_class(r['seq_p'])}'>{fmt_p(r['seq_p'])}</td>"
                     f"<td>{fmt(f.stats['deviance'], 5)}</td><td>{fmt(f.stats['aic'], 5)}</td>"
                     f"<td>{fmt(f.stats['mcfadden'], 4)}</td><td class='l'>{' '.join(note)}</td></tr>")
        h.append("</table><p class='note'>Sequential tests use the likelihood ratio (deviance difference, χ²).</p>")
        return "".join(h)
    h.append("<table><tr><th class='l'>Source</th><th>Sequential p</th><th>Lack of Fit p</th>"
             "<th>Adjusted R²</th><th>Predicted R²</th><th class='l'>Remarks</th></tr>")
    for r in rows[1:]:
        f = r["fit"]
        note = []
        if r["order"] == suggested:
            note.append("<span class='sig'>Suggested</span>")
        if r["aliased"]:
            note.append("<span class='ns'>Aliased</span>")
        h.append(f"<tr><td class='l'>{models.order_label(r['order'])}</td>"
                 f"<td class='{p_class(r['seq_p'])}'>{fmt_p(r['seq_p'])}</td>"
                 f"<td>{fmt_p(r['lof_p'])}</td>"
                 f"<td>{fmt(f.stats.get('adj_r2', np.nan), 4)}</td>"
                 f"<td>{fmt(f.stats.get('pred_r2', np.nan), 4)}</td>"
                 f"<td class='l'>{' '.join(note)}</td></tr>")
    h.append("</table>")
    h.append("<p class='note'><b>How to read:</b> choose the highest-order model whose sequential p is "
             "&lt; 0.05, whose lack of fit is not significant (p &gt; 0.10), that is not <i>aliased</i>, and "
             "whose Adjusted R² and Predicted R² are high.</p>")
    if suggested == "mean":
        h.append("<p class='ns'>No model is significant compared with the mean.</p>")
    return "".join(h)


def anova_html(project, j, fit, heading="h2"):
    kind = project.analysis_kind(j)
    if kind == "glm":
        return glm_html(project, j, fit, heading)
    if kind == "reml":
        return reml_html(project, j, fit, heading)
    h = [f"<{heading}>ANOVA - {project.response_label(j)}</{heading}>"]
    notes = ["Type III (partial) sum of squares."]
    notes.append("Scheffé model in pseudo-components." if project.is_mixture else "Factors in coded units.")
    tr = project.model_spec(j).get("transform")
    if tr and tr.get("kind", "none") != "none":
        notes.append(f"Response transformation: <b>{models.transform_label(tr)}</b>.")
    if project.n_blocks > 1:
        notes.append("The block effect is removed first and is not tested.")
    h.append(f"<p class='note'>{' '.join(notes)}</p>")
    if fit.aliased:
        h.append("<p class='ns'>Aliased terms (removed automatically): " +
                 ", ".join(models.term_name(t) for t in fit.aliased) + "</p>")
    if not fit.anova:
        h.append("<p class='ns'>Residual degrees of freedom = 0 (saturated model). Statistical tests cannot "
                 "be computed - remove model terms or add runs.</p>")
    else:
        h.append("<table><tr><th class='l'>Source</th><th>Sum of Squares</th><th>df</th>"
                 "<th>Mean Square</th><th>F-value</th><th>p-value</th><th class='l'></th></tr>")
        for a in fit.anova:
            cls = "model" if a.source == "Model" else ""
            verdict = ""
            if not np.isnan(a.p):
                good = a.p >= 0.05 if a.source == "Lack of Fit" else a.p < 0.05
                label = "significant" if a.p < 0.05 else "not significant"
                verdict = f"<span class='{'sig' if good else 'ns'}'>{label}</span>"
            h.append(f"<tr class='{cls}'><td class='l'>{a.source}</td><td>{fmt(a.ss)}</td>"
                     f"<td>{a.df}</td><td>{fmt(a.ms)}</td><td>{fmt(a.f, 4)}</td>"
                     f"<td>{fmt_p(a.p)}</td><td class='l'>{verdict}</td></tr>")
        h.append("</table>")
        model_row = next(a for a in fit.anova if a.source == "Model")
        if model_row.p < 0.05:
            h.append(f"<p>The model F-value of {fmt(model_row.f, 4)} implies the model is "
                     f"<b class='sig'>significant</b> (p {fmt_p(model_row.p)}).</p>")
        else:
            h.append("<p class='ns'>The model is not significant relative to noise.</p>")

    st = fit.stats
    h.append("<h3>Fit Statistics</h3><table>")
    for key, label in STAT_LABELS:
        if key in st:
            h.append(f"<tr><td class='l'>{label}</td><td>{fmt(st[key], 5)}</td></tr>")
    h.append("</table>")
    if "pred_r2" in st and not np.isnan(st["pred_r2"]):
        diff = st["adj_r2"] - st["pred_r2"]
        ok = "are in reasonable agreement (difference &lt; 0.2)" if diff < 0.2 else \
            "<span class='ns'>differ by more than 0.2 - check for outliers or simplify the model</span>"
        h.append(f"<p class='note'>Predicted R² and Adjusted R² {ok}.</p>")
    if "adeq_precision" in st:
        ap = st["adeq_precision"]
        h.append(f"<p class='note'>Adeq Precision = {fmt(ap, 4)} " +
                 ("(&gt; 4, adequate signal - the model can be used to navigate the design space)."
                  if ap > 4 else "<span class='ns'>(≤ 4, weak signal).</span>") + "</p>")

    space = "Pseudo-components" if project.is_mixture else "Coded Units"
    h.append(f"<h3>Coefficients ({space})</h3><table><tr><th class='l'>Term</th>"
             "<th>Estimate</th><th>df</th><th>Std. Error</th><th>95% CI Low</th>"
             "<th>95% CI High</th><th>VIF</th></tr>")
    tcrit = fit.t_crit()
    for lab, b, se, v in zip(fit.coef_labels(), fit.coef, fit.coef_se, fit.coef_vif):
        h.append(f"<tr><td class='l'>{lab}</td><td>{fmt(b)}</td><td>1</td>"
                 f"<td>{fmt(se)}</td><td>{fmt(b - tcrit * se)}</td><td>{fmt(b + tcrit * se)}</td>"
                 f"<td>{'' if np.isnan(v) else fmt(v, 4)}</td></tr>")
    h.append("</table>")
    if project.is_mixture:
        h.append("<p class='note'>In mixture models the linear terms are not tested individually; they are tested "
                 "together in the <i>Linear Mixture</i> row.</p>")
    if project.cat_idx:
        h.append("<p class='note'>Categorical factors use effect coding: coefficient X[level] = difference of that "
                 "level from the mean of all levels; last level = −(sum of the other level coefficients).</p>")
    if "shapiro_p" in st and not np.isnan(st["shapiro_p"]):
        sp_ = st["shapiro_p"]
        h.append(f"<p class='note'>Shapiro-Wilk normality test of residuals: p = {fmt_p(sp_)} "
                 + ("(residuals can be considered normal)." if sp_ >= 0.05
                    else "<span class='ns'>(residuals are not normal - consider a transformation).</span>") + "</p>")
    return "".join(h)


def _coef_table(fit, extra_cols=None, zstat=False):
    lab = "z" if zstat else "t"
    h = ["<table><tr><th class='l'>Term</th><th>Estimate</th><th>Std. Error</th>"
         f"<th>{lab}</th><th>p-value</th><th>95% CI Low</th><th>95% CI High</th>"
         + "".join(f"<th>{c}</th>" for c, _ in (extra_cols or [])) + "</tr>"]
    crit = fit.t_crit()
    for n, (lb, b, se) in enumerate(zip(fit.coef_labels(), fit.coef, fit.coef_se)):
        tv = b / se if se > 0 else np.nan
        pv = 2 * (stats.norm.sf(abs(tv)) if zstat else stats.t.sf(abs(tv), max(fit.df_resid, 1)))
        h.append(f"<tr><td class='l'>{lb}</td><td>{fmt(b)}</td><td>{fmt(se)}</td><td>{fmt(tv, 4)}</td>"
                 f"<td class='{p_class(pv)}'>{fmt_p(pv)}</td><td>{fmt(b - crit * se)}</td><td>{fmt(b + crit * se)}</td>"
                 + "".join(f"<td>{fmt(fn(b), 5)}</td>" for _, fn in (extra_cols or [])) + "</tr>")
    h.append("</table>")
    return "".join(h)


def glm_html(project, j, fit, heading="h2"):
    fam = fit.family
    st = fit.stats
    title = "Logistic Regression" if fam == "binomial" else "Poisson Regression"
    h = [f"<{heading}>{title} - {project.response_label(j)}</{heading}>",
         f"<p class='note'>Link: {'logit' if fam == 'binomial' else 'log'}. "
         + (f"Trials per run = {fmt(fit.trials)}. " if fam == "binomial" else "")
         + "Terms are tested with the likelihood ratio (χ²).</p>"]
    h.append("<table><tr><th class='l'>Source</th><th>χ²</th><th>df</th><th>p-value</th><th class='l'></th></tr>")
    for a in fit.anova:
        if a.source.startswith("Residual"):
            continue
        cls = "model" if a.source == "Model" else ""
        verdict = ""
        if not np.isnan(a.p):
            verdict = f"<span class='{p_class(a.p)}'>{'significant' if a.p < 0.05 else 'not significant'}</span>"
        h.append(f"<tr class='{cls}'><td class='l'>{a.source}</td><td>{fmt(a.f, 5)}</td><td>{a.df}</td>"
                 f"<td>{fmt_p(a.p)}</td><td class='l'>{verdict}</td></tr>")
    h.append("</table>")
    h.append("<h3>Fit Statistics</h3><table>")
    rows = [("Deviance", st["deviance"]), ("Null deviance", st["null_deviance"]),
            ("Log-likelihood", st["loglik"]), ("AIC", st["aic"]), ("BIC", st["bic"]),
            ("McFadden pseudo R²", st["mcfadden"]), ("Adj. McFadden R²", st["adj_mcfadden"])]
    if "tjur" in st:
        rows.append(("Tjur R²", st["tjur"]))
    if "dispersion" in st:
        rows.append(("Dispersion (Pearson χ²/df)", st["dispersion"]))
    for lab, v in rows:
        h.append(f"<tr><td class='l'>{lab}</td><td>{fmt(v, 5)}</td></tr>")
    h.append("</table><h3>Goodness-of-Fit Tests</h3><table><tr><th class='l'>Test</th><th>χ²</th><th>df</th>"
             "<th>p-value</th></tr>")
    if "deviance_p" in st:
        h.append(f"<tr><td class='l'>Deviance</td><td>{fmt(st['deviance'])}</td><td>{fit.df_resid}</td>"
                 f"<td>{fmt_p(st['deviance_p'])}</td></tr>")
        h.append(f"<tr><td class='l'>Pearson</td><td>{fmt(st['pearson_chi2'])}</td><td>{fit.df_resid}</td>"
                 f"<td>{fmt_p(st['pearson_p'])}</td></tr>")
    if "hosmer_lemeshow" in st:
        hl = st["hosmer_lemeshow"]
        h.append(f"<tr><td class='l'>Hosmer-Lemeshow</td><td>{fmt(hl['chi2'])}</td><td>{hl['df']}</td>"
                 f"<td>{fmt_p(hl['p'])}</td></tr>")
    h.append("</table><p class='note'>A goodness-of-fit p-value &gt; 0.05 means there is no evidence of lack of fit. "
             "Dispersion well above 1 indicates overdispersion.</p>")
    space = "Pseudo-components" if project.is_mixture else "Coded Units"
    ratio = ("Odds ratio", np.exp) if fam == "binomial" else ("Rate ratio", np.exp)
    h.append(f"<h3>Coefficients ({space}, {'logit' if fam == 'binomial' else 'log'} scale)</h3>")
    h.append(_coef_table(fit, [ratio], zstat=True))
    return "".join(h)


def reml_html(project, j, fit, heading="h2"):
    st = fit.stats
    h = [f"<{heading}>Split-Plot Analysis (REML) - {project.response_label(j)}</{heading}>",
         "<p class='note'>Mixed model: whole plots as a random effect. Wald F tests; denominator df by the "
         "containment method (hard-to-change factor terms are tested against the whole-plot df).</p>"]
    tr = project.model_spec(j).get("transform")
    if tr and tr.get("kind", "none") != "none":
        h.append(f"<p class='note'>Response transformation: <b>{models.transform_label(tr)}</b>.</p>")
    h.append("<table><tr><th class='l'>Term</th><th>Type</th><th>Numerator df</th><th>Denominator df</th>"
             "<th>F-value</th><th>p-value</th><th class='l'></th></tr>")
    wp_names = {models.term_name(t) for t in getattr(fit, "wp_terms", [])}
    for a in fit.anova:
        verdict = f"<span class='{p_class(a.p)}'>{'significant' if a.p < 0.05 else 'not significant'}</span>"
        h.append(f"<tr><td class='l'>{a.source}</td><td>{'Whole plot' if a.source in wp_names else 'Subplot'}</td>"
                 f"<td>{a.df}</td><td>{int(a.ms)}</td><td>{fmt(a.f, 4)}</td><td>{fmt_p(a.p)}</td>"
                 f"<td class='l'>{verdict}</td></tr>")
    h.append("</table><h3>Variance Components</h3><table><tr><th class='l'>Source</th><th>Variance</th>"
             "<th>Std. Dev.</th><th>% total</th></tr>")
    tot = st["sigma2"] + st["sigma2_wp"]
    for lab, v in (("Whole plot (between groups)", st["sigma2_wp"]), ("Subplot (residual)", st["sigma2"])):
        h.append(f"<tr><td class='l'>{lab}</td><td>{fmt(v)}</td><td>{fmt(np.sqrt(v))}</td>"
                 f"<td>{fmt(100 * v / tot if tot else np.nan, 4)}</td></tr>")
    h.append("</table><h3>Statistics</h3><table>")
    for lab, v in (("Number of whole plots", st["n_wp"]), ("df whole plot", st["df_wp"]), ("df subplot", st["df_sp"]),
                   ("Variance ratio γ = σ²wp/σ²", st["gamma"]), ("-2 REML log-likelihood", st["reml_m2ll"]),
                   ("Conditional R²", st["r2_cond"]), ("Mean", st["mean"])):
        h.append(f"<tr><td class='l'>{lab}</td><td>{fmt(v, 5)}</td></tr>")
    h.append("</table>")
    if st["df_wp"] <= 0:
        h.append("<p class='ns'>Whole-plot df = 0: hard-to-change factor effects cannot be tested properly. "
                 "Add more whole plots (replicates).</p>")
    h.append("<h3>Coefficients (GLS)</h3>")
    h.append(_coef_table(fit))
    return "".join(h)


def equation_html(project, j, fit, heading="h2"):
    lhs = response_lhs(project, j)
    names = [f.name for f in project.factors]
    h = [f"<{heading}>Model Equation - {project.response_label(j)}</{heading}>"]
    if project.is_mixture:
        pseudo = models.format_equation(lhs, fit.col_terms, fit.coef, lhs=lhs)
        h += ["<h3>In Pseudo-components</h3>", f"<pre class='box'>{pseudo}</pre>",
              "<p class='note'>Pseudo-component: x = (amount − lower limit) / (total − Σ lower limits). "
              "Linear coefficients = predicted response of the pure blends (vertices).</p>"]
        act = None if project.is_combined else \
            models.mixture_actual_coefficients(fit.col_terms, fit.coef, project.lows, project.mixture_total)
        if act is not None:
            eq = models.format_equation(lhs, act[0], act[1], names=names, lhs=lhs)
            h += ["<h3>In Actual Components</h3>", f"<pre class='box'>{eq}</pre>"]
        else:
            h.append("<p class='note'>The equation in actual components is shown only when all lower "
                     "limits = 0. Use the pseudo-component equation above.</p>")
        legend = "<br>".join(f"{project.factor_label(i)}: {fmt(f.low)} – {fmt(f.high)}"
                             for i, f in enumerate(project.factors))
        legend += f"<br>Mixture total = {fmt(project.mixture_total)}"
    else:
        coded_eq = models.format_items(lhs, list(zip(fit.coef_labels(), fit.coef)))
        h += ["<h3>In Coded Factors</h3>", f"<pre class='box'>{coded_eq}</pre>",
              "<p class='note'>The coded equation is used to compare the relative effect of each factor "
              "(high level = +1, low = -1).</p>"]
        for combo, at, ab in models.actual_equations(fit, project.lows, project.highs):
            title = "In Actual Factors"
            if combo:
                title += " - " + ", ".join(f"{project.factors[i].name} = {project.factors[i].levels[lv]}"
                                           for i, lv in combo.items())
            eq = models.format_equation(lhs, at, ab, names=names, lhs=lhs)
            h += [f"<h3>{title}</h3>", f"<pre class='box'>{eq}</pre>"]
        h.append("<p class='note'>The actual equation is used for predictions in original units. Do not use it "
                 "to judge relative effects because its coefficients depend on the scale of the units.</p>")
        legend = "<br>".join(
            f"{project.factor_label(i)}: level = {', '.join(f.levels)}" if f.categoric else
            f"{project.factor_label(i)}: low = {fmt(f.low)}, high = {fmt(f.high)}"
            for i, f in enumerate(project.factors))
    if project.n_blocks > 1:
        legend += "<br>The equation applies to the average of all blocks."
    h.append(f"<p class='note'>{legend}</p>")
    return "".join(h)


# ------------------------------------------------------------------ design & optimization HTML
def design_html(project):
    p = project
    unit = "Component" if p.is_mixture else "Factor"
    h = [f"<h2>Design Summary</h2><p><b>{p.design_description()}</b></p>"]
    h.append(f"<table><tr><th class='l'>{unit}</th><th class='l'>Name</th><th class='l'>Unit</th>"
             + ("<th>Lower Limit</th><th>Upper Limit</th>" if p.is_mixture else "<th>Low (-1)</th><th>High (+1)</th>")
             + "</tr>")
    for i, f in enumerate(p.factors):
        rng = f"<td colspan='2' class='l'>Categorical: {', '.join(f.levels)}</td>" if f.categoric else \
            f"<td>{fmt(f.low)}</td><td>{fmt(f.high)}</td>"
        h.append(f"<tr><td class='l'>{designs.LETTERS[i]}</td><td class='l'>{f.name}</td>"
                 f"<td class='l'>{f.unit}</td>{rng}</tr>")
    h.append("</table>")
    if p.constraints:
        h.append("<p><b>Constraint:</b> " + "; ".join(c["text"] for c in p.constraints) + "</p>")
    h.append("<h3>Data Sheet</h3><table><tr><th>Std</th><th>Run</th>"
             + ("<th>Block</th>" if p.n_blocks > 1 else "") + "<th class='l'>Type</th>"
             + "".join(f"<th>{designs.LETTERS[i]}: {f.name}</th>" for i, f in enumerate(p.factors))
             + "".join(f"<th>R{j + 1}: {r.name}</th>" for j, r in enumerate(p.responses)) + "</tr>")
    actual = p.actual
    order = p.display_order()
    for idx in order[:PDF_MAX_ROWS]:
        h.append(f"<tr><td>{idx + 1}</td><td>{p.run_order[idx]}</td>"
                 + (f"<td>{p.blocks[idx]}</td>" if p.n_blocks > 1 else "")
                 + f"<td class='l'>{p.point_types[idx]}</td>"
                 + "".join(f"<td>{p.format_value(i, v)}</td>" for i, v in enumerate(actual[idx]))
                 + "".join(f"<td>{'' if np.isnan(v) else fmt(v, 6)}</td>" for v in p.data[idx]) + "</tr>")
    h.append("</table>")
    if len(order) > PDF_MAX_ROWS:
        h.append(f"<p class='note'>Showing {PDF_MAX_ROWS} of {len(order)} runs. The complete data are in the Excel "
                 "export.</p>")
    return "".join(h)


def optimization_html(project, criteria, solutions, max_rows=10):
    from doe.optimize import GOALS
    p = project
    h = ["<h2>Numerical Optimization (Desirability)</h2><h3>Criteria</h3>",
         "<table><tr><th class='l'>Name</th><th class='l'>Goal</th><th>Lower Limit</th><th>Upper Limit</th>"
         "<th>Target</th><th>Importance</th></tr>"]
    for i, f in enumerate(p.factors):
        if f.categoric:
            h.append(f"<tr><td class='l'>{p.factor_label(i)}</td><td class='l'>All levels</td>"
                     "<td></td><td></td><td></td><td></td></tr>")
            continue
        c = p.opt_criteria.get(f"f{i}", {"low": f.low, "high": f.high})
        h.append(f"<tr><td class='l'>{p.factor_label(i)}</td><td class='l'>In range</td>"
                 f"<td>{fmt(c['low'])}</td><td>{fmt(c['high'])}</td><td></td><td></td></tr>")
    if p.constraints:
        h.append("<tr><td class='l' colspan='6'>Constraint: " + "; ".join(c["text"] for c in p.constraints)
                 + "</td></tr>")
    for c in criteria:
        target = fmt(c.get("target")) if c["goal"] == "target" else ""
        h.append(f"<tr><td class='l'>{c.get('label') or p.response_label(c['index'])}</td><td class='l'>{GOALS[c['goal']]}</td>"
                 f"<td>{fmt(c['low'])}</td><td>{fmt(c['high'])}</td><td>{target}</td>"
                 f"<td>{'+' * int(c.get('importance', 3))}</td></tr>")
    h.append("</table><h3>Solutions</h3><table><tr><th>No</th>"
             + "".join(f"<th>{f.name}</th>" for f in p.factors)
             + "".join(f"<th>{r.name}</th>" for r in p.responses) + "<th>Desirability</th></tr>")
    fits = [p.model(j) for j in range(len(p.responses))]
    for n, (D, x) in enumerate(solutions[:max_rows]):
        act = p.to_actual(x)
        preds = [fmt(float(f.predict(x[None, :])[0]), 5) if f else "-" for f in fits]
        h.append(f"<tr><td>{n + 1}</td>" + "".join(f"<td>{p.format_value(i, v, 5)}</td>" for i, v in enumerate(act))
                 + "".join(f"<td>{v}</td>" for v in preds) + f"<td><b>{D:.3f}</b></td></tr>")
    h.append("</table>")
    return "".join(h)


# ------------------------------------------------------------------ PDF
def _figure_png(draw_fn, size=(9, 6)):
    fig = Figure(figsize=size, dpi=110, layout="constrained")
    FigureCanvasAgg(fig)
    draw_fn(fig)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=110)
    return buf.getvalue()


def default_graph(project, fit, j):
    """Main graph for the report: contour/ternary plot of the first two factors."""
    label = project.response_label(j)
    if project.is_mixture:
        comps = [0, 1, 2]
        center = np.nanmean(project.coded, axis=0)
        fixed = {i: float(center[i]) for i in range(3, project.k)}
        return lambda fig: plots.draw_ternary(fig, project, fit, comps, fixed, label)
    num = project.num_idx
    base = np.zeros(project.k)
    if len(num) >= 2:
        return lambda fig: plots.draw_process_2d(fig, project, fit, num[0], num[1], base, label)
    return lambda fig: plots.draw_oneway(fig, project, fit, num[0] if num else 0, base, label)


PDF_MAX_ROWS = 200


def write_pdf(project, path, criteria=None, solutions=None, app_name="REGEN"):
    from PySide6.QtCore import QMarginsF, QUrl
    from PySide6.QtGui import QImage, QPageLayout, QPageSize, QPdfWriter, QTextDocument

    doc = QTextDocument()
    images = []

    def add_image(png, width=640):
        name = f"img{len(images)}.png"
        img = QImage.fromData(png, "PNG")
        doc.addResource(QTextDocument.ImageResource, QUrl(name), img)
        images.append(img)
        return f"<p><img src='{name}' width='{width}'></p>"

    p = project
    title = p.path.replace("\\", "/").split("/")[-1] if p.path else "Untitled project"
    parts = [REPORT_CSS, f"<h1 style='color:#1f1f1f'>Experiment Analysis Report</h1>",
             f"<p class='note'>{title} - created with {app_name}, "
             f"{datetime.datetime.now().strftime('%d-%m-%Y %H:%M')}</p>"]
    if (p.notes or "").strip():
        parts.append("<h2>Notes</h2><div class='box'><pre style='white-space: pre-wrap'>"
                     + p.notes.replace("&", "&amp;").replace("<", "&lt;") + "</pre></div>")
    parts.append(design_html(p))

    for j in range(len(p.responses)):
        fit, err = p.fit_with_error(j)
        parts.append(f"<h1 style='color:#1f1f1f; page-break-before: always'>{p.response_label(j)}</h1>")
        if fit is None:
            parts.append(f"<p class='ns'>{err}</p>")
            continue
        spec = p.model_spec(j)
        parts.append(f"<p>Model: <b>{models.order_label(spec['order'])}</b>"
                     + (" (terms selected manually / by elimination)" if spec["terms"] is not None else "")
                     + f", transformation: <b>{models.transform_label(spec.get('transform'))}</b></p>")
        rows, suggested = p.fit_summary(j)
        parts.append(summary_html(p, j, rows, suggested, heading="h2"))
        parts.append(anova_html(p, j, fit, heading="h2"))
        parts.append(equation_html(p, j, fit, heading="h2"))
        run = p.run_order[fit.rows_used]
        parts.append("<h2>Diagnostics</h2>")
        parts.append(add_image(_figure_png(lambda fig: plots.draw_diagnostics(fig, fit, run), (11, 6.5))))
        parts.append("<h2>Model Graphs</h2>")
        parts.append(add_image(_figure_png(default_graph(p, fit, j), (8, 6)), width=560))

    if len(p.responses) > 1:
        from .info_tabs import coef_table_html
        parts.append("<div style='page-break-before: always'></div>" + coef_table_html(p))

    from doe import adequacy

    from .ann_tab import (adequacy_html, ann_summary_html, draw_compare, draw_importance, draw_network,
                          draw_regression, weights_html)
    for j, m in sorted(p.ann.items()):
        parts.append(f"<h1 style='color:#1f1f1f; page-break-before: always'>ANN - {p.response_label(j)}</h1>")
        parts.append(f"<p>Model used for graphs/optimization: <b>{'ANN' if p.source(j) == 'ann' else 'RSM'}"
                     "</b></p>")
        try:
            parts.append(adequacy_html(p, j, adequacy.assess(p, j)).replace("<h2>", "<h2 style='margin-top:0'>", 1))
        except Exception:  # noqa: BLE001 - adequacy assessment is optional in the report
            pass
        parts.append(ann_summary_html(p, j, m))
        parts.append(add_image(_figure_png(lambda fig, m=m: draw_regression(fig, m), (9, 7))))
        parts.append(add_image(_figure_png(lambda fig, j=j, m=m: draw_compare(fig, p, j, m), (10, 4.5))))
        parts.append(add_image(_figure_png(lambda fig, m=m: draw_importance(fig, p, m), (7, 4)), width=480))
        parts.append(add_image(_figure_png(lambda fig, j=j, m=m: draw_network(fig, p, m, p.responses[j].name), (9, 6)), width=560))
        parts.append(weights_html(p, m))
    if isinstance(p.learning, dict) and p.learning:
        from .ann_tab import draw_learning, learning_html
        for j, lc in sorted(p.learning.items()):
            parts.append("<div style='page-break-before: always'></div>" + learning_html(p, j, lc))
            parts.append(add_image(_figure_png(lambda fig, j=j, lc=lc: draw_learning(fig, p, j, lc), (9, 4.5)),
                                   width=600))

    try:
        from doe import readiness

        from .readiness_tab import readiness_html
        for j in range(len(p.responses)):
            a = readiness.assess(p, j)
            head = "<h1 style='color:#1f1f1f; page-break-before: always'>Data Readiness & Recommendations</h1>"                 if j == 0 else "<div style='page-break-before: always'></div>"
            parts.append(head + readiness_html(p, j, a))
    except Exception as exc:  # noqa: BLE001 - optional report section
        parts.append(f"<p class='ns'>The data readiness assessment could not be created: {exc}</p>")

    if criteria and solutions:
        parts.append("<div style='page-break-before: always'></div>")
        parts.append(optimization_html(p, criteria, solutions))

    res = p.nsga_result
    if res:
        from .nsga_tab import ALG_NAME, draw_pareto, objective_label
        alg = ALG_NAME.get(res.get("algorithm"), "NSGA-II")
        parts.append(f"<h1 style='color:#1f1f1f; page-break-before: always'>{alg} Multi-objective Optimization</h1>")
        prm = res["params"]
        parts.append(f"<p>Population {prm['pop']}, maximum generations {prm['gens']} (ran "
                     f"{res.get('gens_run', prm['gens'])}"
                     + (f"; {res['stop_reason']}" if res.get("stop_reason") else "") + f"), P crossover {prm['pc']}, "
                     f"η crossover {prm['eta_c']}, η mutation {prm['eta_m']}, seed {prm['seed']}. "
                     f"{len(res['X'])} Pareto solutions.</p>")
        chk = res.get("check") or {}
        if chk.get("notes"):
            parts.append("<ul>" + "".join(f"<li>{n}</li>" for n in chk["notes"]) + "</ul>")
        parts.append("<table><tr><th class='l'>Objective</th><th class='l'>Goal</th><th>Lower limit</th>"
                     "<th>Upper limit</th><th>TOPSIS weight</th></tr>")
        for o in res["objectives"]:
            parts.append(f"<tr><td class='l'>{objective_label(p, o)}</td><td class='l'>{o['goal']}</td>"
                         f"<td>{fmt(o.get('low'))}</td><td>{fmt(o.get('high'))}</td><td>{fmt(o.get('weight'))}</td></tr>")
        parts.append("</table>")
        m_obj = len(res["active"])
        parts.append(add_image(_figure_png(lambda fig: draw_pareto(
            fig, p, res, None, "front" if m_obj == 2 else "front2d",
            {"x": 0, "y": 1, "color": 2 if m_obj > 2 else "topsis", "size": 3 if m_obj > 3 else None}), (8, 6)),
            width=560))
        if m_obj > 3:
            parts.append(add_image(_figure_png(lambda fig: draw_pareto(fig, p, res, None, "parallel"), (9, 5)),
                                   width=600))
        if res.get("hv_curve"):
            parts.append(add_image(_figure_png(lambda fig: draw_pareto(fig, p, res, None, "convergence"), (8, 5.5)),
                                   width=520))
        X = p.to_actual(res["X"])
        parts.append("<h3>Pareto Solutions (10 Highest TOPSIS Scores)</h3><table><tr><th>No</th>"
                     + "".join(f"<th>{f.name}</th>" for f in p.factors)
                     + "".join(f"<th>{r.name}</th>" for r in p.responses) + "<th>TOPSIS</th></tr>")
        for r in np.argsort(-res["topsis"])[:10]:
            mark = " ★" if r == res["i_topsis"] else ""
            parts.append(f"<tr><td>{r + 1}{mark}</td>" + "".join(f"<td>{p.format_value(i, X[r, i], 5)}</td>"
                                                               for i in range(p.k))
                         + "".join(f"<td>{fmt(float(res['Y'][j][r]), 5) if j in res['Y'] else '-'}</td>"
                                   for j in range(len(p.responses)))
                         + f"<td><b>{res['topsis'][r]:.4f}</b></td></tr>")
        parts.append("</table>")
    cmp = getattr(p, "nsga_compare", None)
    if cmp:
        from .nsga_tab import compare_html, draw_compare
        parts.append("<div style='page-break-before: always'></div>")
        parts.append(compare_html(p, cmp))
        parts.append(add_image(_figure_png(lambda fig: draw_compare(fig, p, cmp, "cmp_front", {"x": 0, "y": 1}),
                                           (8, 6)), width=540))
        parts.append(add_image(_figure_png(lambda fig: draw_compare(fig, p, cmp, "cmp_cross"), (9, 5)), width=600))

    doc.setHtml("".join(parts))
    writer = QPdfWriter(path)
    writer.setPageSize(QPageSize(QPageSize.A4))
    writer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    writer.setResolution(96)
    writer.setTitle(f"Report {title}")
    writer.setCreator(app_name)
    doc.setPageSize(writer.pageLayout().paintRectPixels(96).size().toSizeF())
    doc.print_(writer)


# ------------------------------------------------------------------ Excel
def write_excel(p, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Design"
    head_fill = PatternFill("solid", fgColor="E8EEF6")
    ws.append([p.design_description()])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    headers = ["Std", "Run"] + (["Block"] if p.n_blocks > 1 else []) + ["Type"] + \
              [p.factor_label(i) for i in range(p.k)] + [p.response_label(j) for j in range(len(p.responses))]
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    actual = p.actual
    for idx in p.display_order():
        ws.append([int(idx + 1), int(p.run_order[idx])] + ([int(p.blocks[idx])] if p.n_blocks > 1 else [])
                  + [p.point_types[idx]]
                  + [p.format_value(i, v) if p.factors[i].categoric else (None if np.isnan(v) else float(v))
                     for i, v in enumerate(actual[idx])]
                  + [None if np.isnan(v) else float(v) for v in p.data[idx]])
    for col in ws.columns:
        ws.column_dimensions[col[2].column_letter].width = 16

    def bold_row(sh):
        for c in sh[sh.max_row]:
            c.font = Font(bold=True)
            c.fill = head_fill

    names = [f.name for f in p.factors]
    for j, r in enumerate(p.responses):
        fit = p.fit(j)
        if fit is None:
            continue
        sh = wb.create_sheet(f"R{j + 1} {r.name}"[:31])
        sh.append([f"Analysis of {p.response_label(j)}"])
        sh["A1"].font = Font(bold=True, size=12)
        sh.append([f"Transformation: {models.transform_label(p.model_spec(j).get('transform'))}"])
        sh.append([])
        sh.append(["ANOVA", "Sum of Squares", "df", "Mean Square", "F-value", "p-value"])
        bold_row(sh)
        for a in fit.anova:
            sh.append([a.source, a.ss, a.df] + [None if np.isnan(v) else float(v) for v in (a.ms, a.f, a.p)])
        sh.append([])
        sh.append(["Statistic", "Value"])
        bold_row(sh)
        for key, label in STAT_LABELS:
            v = fit.stats.get(key)
            if v is not None and not np.isnan(v):
                sh.append([label, float(v)])
        sh.append([])
        space = "pseudo-component" if p.is_mixture else "coded"
        sh.append(["Term", f"Coefficient ({space})", "Std. Error"])
        bold_row(sh)
        for lab, b, se in zip(fit.coef_labels(), fit.coef, fit.coef_se):
            sh.append([lab, float(b), None if np.isnan(se) else float(se)])
        sh.append([])
        if p.is_mixture:
            act = None if p.is_combined else \
                models.mixture_actual_coefficients(fit.col_terms, fit.coef, p.lows, p.mixture_total)
            eqs = [({}, act[0], act[1])] if act is not None else []
        else:
            eqs = models.actual_equations(fit, p.lows, p.highs)
        for combo, at, ab in eqs:
            title = "Term (actual)" + ("" if not combo else " - " + ", ".join(
                f"{p.factors[i].name}={p.factors[i].levels[lv]}" for i, lv in combo.items()))
            sh.append([title, "Actual coefficient"])
            bold_row(sh)
            for t, b in zip(at, ab):
                sh.append([models.term_name(t, names), float(b)])
            sh.append([])
        sh.column_dimensions["A"].width = 28
        for col in "BCDEF":
            sh.column_dimensions[col].width = 16
    wb.save(path)
