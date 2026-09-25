"""Import an existing data set (paste from Excel or another DOE program, or open an Excel/CSV file)."""
import csv
import io
import math
import re

import numpy as np
from PySide6.QtCore import QAbstractTableModel, QModelIndex, Qt
from PySide6.QtGui import QAction, QBrush, QColor, QGuiApplication
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog,
                               QFormLayout, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMenu,
                               QMessageBox, QPushButton, QSplitter, QTableView, QTableWidget, QVBoxLayout,
                               QWidget)

from doe.designs import LETTERS
from doe.project import Factor, Project, Response

from .common import fmt, parse_float

ROLES = [("factor", "Numeric factor"), ("categoric", "Categoric factor"), ("component", "Mixture component"),
         ("response", "Response"), ("ignore", "Ignore")]
ROLE_LABEL = dict(ROLES)
ROLE_COLOR = {"factor": "#e3edf9", "categoric": "#f3e8fd", "component": "#e6f4ea", "response": "#fff4d6",
              "ignore": "#eeeeee"}
# header names of bookkeeping columns (English, Indonesian, and the style of other DOE programs)
IGNORE_NAMES = re.compile(r"^(std|run|no\.?|nomor|number|order|urutan|id|block|blok|#)$", re.I)


def _num(text):
    try:
        v = parse_float(text)
    except ValueError:
        return None
    return v


def parse_text(text):
    """Clipboard/CSV text -> list of rows. Separator: tab (Excel), semicolon, or comma."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    if not text.strip():
        return []
    lines = text.split("\n")
    if any("\t" in ln for ln in lines):
        return [ln.split("\t") for ln in lines]
    sample = "\n".join(lines[:20])
    delim = ";" if sample.count(";") >= sample.count(",") and ";" in sample else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def read_file(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (repr(v) if isinstance(v, float) else str(v)) for v in r])
        while rows and not any(c.strip() for c in rows[-1]):
            rows.pop()
        return rows
    with open(path, "r", encoding="utf-8-sig", errors="replace") as fh:
        return parse_text(fh.read())


class ImportModel(QAbstractTableModel):
    """Import preview table as a virtual model (light enough for tens of thousands of rows)."""

    def __init__(self, dlg):
        super().__init__(dlg)
        self.dlg = dlg

    def reset(self):
        self.beginResetModel()
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):  # noqa: N802
        return 0 if parent.isValid() else len(self.dlg.raw)

    def columnCount(self, parent=QModelIndex()):  # noqa: N802
        return 0 if parent.isValid() else len(self.dlg.cols)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        v = self.dlg.raw[index.row()][index.column()]
        crole = self.dlg.cols[index.column()]["role"]
        if role == Qt.DisplayRole:
            return v
        if role == Qt.BackgroundRole:
            bad = crole in ("factor", "component", "response") and v != "" and _num(v) is None
            return QBrush(QColor("#fde2e1" if bad else ROLE_COLOR[crole]))
        if role == Qt.ForegroundRole and crole == "ignore":
            return QBrush(QColor("#999"))
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):  # noqa: N802
        if orientation == Qt.Horizontal and section < len(self.dlg.cols):
            c = self.dlg.cols[section]
            if role == Qt.DisplayRole:
                unit = f" ({c['unit']})" if c["unit"] else ""
                return f"{c['name']}{unit}\n[{ROLE_LABEL[c['role']]}]"
            if role == Qt.BackgroundRole:
                return QBrush(QColor(ROLE_COLOR[c["role"]]))
        elif orientation == Qt.Vertical and role == Qt.DisplayRole:
            return str(section + 1)
        return None


class ImportDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Import Data Set")
        self.resize(1080, 700)
        self.project = None
        self.cols = []        # per column: {"name", "unit", "role"}
        self.raw = []         # data rows (without header)

        root = QVBoxLayout(self)
        info = QLabel("<b>Import existing data.</b> 1) Copy the data from Excel or another DOE program. "
                      "2) Click <b>Paste from Clipboard</b> (or <b>Open File</b>). 3) Right-click rows to delete "
                      "rows you do not need. 4) Right-click a <b>column header</b> to mark it as Factor / Response "
                      "(or select a column and set it in the right panel).")
        info.setWordWrap(True)
        info.setStyleSheet("background: #e3f2fb; padding: 8px; border-radius: 4px;")
        root.addWidget(info)

        bar = QHBoxLayout()
        self.btn_paste = QPushButton("Paste from Clipboard")
        self.btn_paste.setStyleSheet("font-weight: bold; padding: 5px 12px;")
        self.btn_file = QPushButton("Open File (Excel/CSV)...")
        self.chk_header = QCheckBox("First row contains column headers")
        self.chk_header.setChecked(True)
        bar.addWidget(self.btn_paste)
        bar.addWidget(self.btn_file)
        bar.addWidget(self.chk_header)
        bar.addStretch()
        root.addLayout(bar)

        split = QSplitter(Qt.Horizontal)
        self.tbl = QTableView()
        self.model = ImportModel(self)
        self.tbl.setModel(self.model)
        self.tbl.horizontalHeader().setResizeContentsPrecision(60)
        self.tbl.setSelectionBehavior(QTableWidget.SelectItems)
        self.tbl.horizontalHeader().setContextMenuPolicy(Qt.CustomContextMenu)
        self.tbl.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tbl.verticalHeader().setContextMenuPolicy(Qt.CustomContextMenu)
        split.addWidget(self.tbl)

        side = QWidget()
        sl = QVBoxLayout(side)
        g = QGroupBox("Selected Column")
        f = QFormLayout(g)
        self.ed_name = QLineEdit()
        self.ed_unit = QLineEdit()
        self.cb_role = QComboBox()
        for key, label in ROLES:
            self.cb_role.addItem(label, key)
        self.lbl_col = QLabel()
        self.lbl_col.setWordWrap(True)
        self.lbl_col.setStyleSheet("color: #555;")
        f.addRow("Name:", self.ed_name)
        f.addRow("Units:", self.ed_unit)
        f.addRow("Role:", self.cb_role)
        f.addRow(self.lbl_col)
        sl.addWidget(g)
        g2 = QGroupBox("Mixture")
        f2 = QFormLayout(g2)
        self.sp_total = QDoubleSpinBox()
        self.sp_total.setRange(0.0001, 1e9)
        self.sp_total.setDecimals(4)
        f2.addRow("Mixture total:", self.sp_total)
        self.g_mix = g2
        sl.addWidget(g2)
        self.lbl_sum = QLabel()
        self.lbl_sum.setWordWrap(True)
        self.lbl_sum.setStyleSheet("font-weight: bold; color: #1f5f99;")
        sl.addWidget(self.lbl_sum)
        sl.addStretch()
        split.addWidget(side)
        split.setSizes([760, 300])
        root.addWidget(split, 1)

        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Create Project")
        bb.button(QDialogButtonBox.Cancel).setText("Cancel")
        bb.accepted.connect(self.create)
        bb.rejected.connect(self.reject)
        root.addWidget(bb)

        self.btn_paste.clicked.connect(self.paste)
        self.btn_file.clicked.connect(self.open_file)
        self.chk_header.toggled.connect(self.reload)
        self.tbl.horizontalHeader().customContextMenuRequested.connect(self.header_menu)
        self.tbl.horizontalHeader().sectionClicked.connect(self.select_column)
        self.tbl.customContextMenuRequested.connect(self.row_menu)
        self.tbl.verticalHeader().customContextMenuRequested.connect(self.row_menu)
        self.tbl.selectionModel().currentChanged.connect(lambda cur, _prev: self.select_column(cur.column()))
        self.ed_name.editingFinished.connect(self.apply_props)
        self.ed_unit.editingFinished.connect(self.apply_props)
        self.cb_role.currentIndexChanged.connect(self.apply_props)
        self._all_rows = []
        self._cur = -1
        self._busy = False
        self.update_summary()

    # --------------------------------------------------------------- loading
    def paste(self):
        rows = parse_text(QGuiApplication.clipboard().text())
        if not rows:
            QMessageBox.information(self, "Import", "The clipboard is empty. Copy a data table first.")
            return
        self.load_rows(rows)

    def open_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Data", "", "Data (*.xlsx *.xlsm *.csv *.txt);;All Files (*.*)")
        if not path:
            return
        try:
            rows = read_file(path)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "Import", f"The file could not be read:\n{exc}")
            return
        self.load_rows(rows)

    def load_rows(self, rows):
        width = max(len(r) for r in rows)
        rows = [[c.strip() for c in r] + [""] * (width - len(r)) for r in rows]
        rows = [r for r in rows if any(r)]
        # drop columns that are completely empty
        keep = [j for j in range(width) if any(r[j] for r in rows)]
        self._all_rows = [[r[j] for j in keep] for r in rows]
        first = self._all_rows[0] if self._all_rows else []
        has_header = any(c and _num(c) is None for c in first) and len(self._all_rows) > 1
        self._busy = True
        self.chk_header.setChecked(has_header)
        self._busy = False
        self.reload()

    def reload(self):
        if self._busy or not self._all_rows:
            return
        header = self.chk_header.isChecked()
        head = self._all_rows[0] if header else []
        self.raw = [list(r) for r in (self._all_rows[1:] if header else self._all_rows)]
        ncol = len(self._all_rows[0])
        self.cols = []
        for j in range(ncol):
            name = head[j] if header and head[j] else f"Column {j + 1}"
            de_response = bool(re.match(r"^R\d+\s*:", name))      # response header "R1: Name (units)" used by other DOE programs
            unit = ""
            m = re.match(r"^(?:[A-Z]:\s*|R\d+:\s*)?(.*?)\s*[\(\[]([^\)\]]*)[\)\]]\s*$", name)
            if m:
                name, unit = m.group(1).strip() or name, m.group(2).strip()
            else:
                name = re.sub(r"^(?:[A-Z]|R\d+):\s*", "", name)
            role = self._guess(j, name)
            if de_response and role == "factor":
                role = "response"
            self.cols.append({"name": name, "unit": unit, "role": role})
        nums = [j for j, c in enumerate(self.cols) if c["role"] == "factor"]
        if nums and not any(c["role"] == "response" for c in self.cols):
            self.cols[nums[-1]]["role"] = "response"
        self.refresh_table()
        self.select_column(0)

    def _values(self, j):
        return [r[j] for r in self.raw if r[j] != ""]

    def _guess(self, j, name):
        vals = self._values(j)
        if IGNORE_NAMES.match(name.strip()) or not vals:
            return "ignore"
        nums = [_num(v) for v in vals]
        numeric_share = sum(v is not None for v in nums) / len(nums)
        # columns that are almost all numbers are treated as numeric; mistyped cells are marked red
        return "factor" if numeric_share >= 0.8 else "categoric"

    # --------------------------------------------------------------- display
    def refresh_table(self):
        self._busy = True
        self.model.reset()
        self.tbl.resizeColumnsToContents()
        self._busy = False
        self.update_summary()

    def select_column(self, j):
        if j is None or j < 0 or j >= len(self.cols):
            return
        self._cur = j
        c = self.cols[j]
        self._busy = True
        self.ed_name.setText(c["name"])
        self.ed_unit.setText(c["unit"])
        self.cb_role.setCurrentIndex(self.cb_role.findData(c["role"]))
        self._busy = False
        vals = self._values(j)
        nums = [v for v in (_num(x) for x in vals) if v is not None]
        if c["role"] == "categoric":
            levels = list(dict.fromkeys(vals))
            self.lbl_col.setText(f"{len(levels)} levels: {', '.join(levels[:12])}{' …' if len(levels) > 12 else ''}")
        elif nums:
            self.lbl_col.setText(f"{len(nums)} values, min {fmt(min(nums))}, max {fmt(max(nums))}, "
                                 f"{len(set(nums))} unique values.")
        else:
            self.lbl_col.setText("")

    def apply_props(self):
        if self._busy or self._cur < 0:
            return
        c = self.cols[self._cur]
        c["name"] = self.ed_name.text().strip() or c["name"]
        c["unit"] = self.ed_unit.text().strip()
        c["role"] = self.cb_role.currentData()
        cur = self._cur
        self.refresh_table()
        self.select_column(cur)

    def header_menu(self, pos):
        j = self.tbl.horizontalHeader().logicalIndexAt(pos)
        if j < 0:
            return
        m = QMenu(self)
        for key, label in ROLES:
            act = QAction(label, m)
            act.setCheckable(True)
            act.setChecked(self.cols[j]["role"] == key)
            act.triggered.connect(lambda _=False, j=j, key=key: self.set_role(j, key))
            m.addAction(act)
        m.exec(self.tbl.horizontalHeader().mapToGlobal(pos))

    def set_role(self, j, role):
        self.cols[j]["role"] = role
        self.refresh_table()
        self.select_column(j)

    def row_menu(self, pos):
        rows = sorted({i.row() for i in self.tbl.selectionModel().selectedIndexes()})
        if not rows:
            return
        m = QMenu(self)
        act = QAction(f"Delete {len(rows)} Selected Rows", m)
        act.triggered.connect(lambda: self.delete_rows(rows))
        m.addAction(act)
        m.exec(self.tbl.viewport().mapToGlobal(pos))

    def delete_rows(self, rows):
        for i in sorted(rows, reverse=True):
            del self.raw[i]
        self.refresh_table()

    def update_summary(self):
        roles = [c["role"] for c in self.cols]
        nf = roles.count("factor") + roles.count("categoric")
        nc = roles.count("component")
        nr = roles.count("response")
        self.g_mix.setVisible(nc > 0)
        if nc:
            idx = [j for j, c in enumerate(self.cols) if c["role"] == "component"]
            sums = [sum(_num(r[j]) or 0 for j in idx) for r in self.raw if all(_num(r[j]) is not None for j in idx)]
            if sums and not self._busy:
                self._busy = True
                self.sp_total.setValue(float(np.median(sums)))
                self._busy = False
        parts = []
        if nc:
            parts.append(f"{nc} mixture components")
        if nf:
            parts.append(f"{nf} factors")
        parts.append(f"{nr} responses")
        self.lbl_sum.setText(", ".join(parts) + f", {len(self.raw)} rows." if self.cols else
                             "No data yet. Paste from the clipboard or open a file.")

    # --------------------------------------------------------------- create project
    def create(self):
        if not self.cols or not self.raw:
            QMessageBox.warning(self, "Import", "No data yet.")
            return
        roles = [c["role"] for c in self.cols]
        f_idx = [j for j, r in enumerate(roles) if r in ("factor", "categoric", "component")]
        r_idx = [j for j, r in enumerate(roles) if r == "response"]
        comp = [j for j in f_idx if roles[j] == "component"]
        if not f_idx or not r_idx:
            QMessageBox.warning(self, "Import", "Mark at least one Factor column and one Response column.")
            return
        if comp and len(comp) != len(f_idx):
            QMessageBox.warning(self, "Import", "Mixture components cannot be combined with process factors when "
                                                 "importing data. Mark all factors as components, or none.")
            return
        if comp and len(comp) < 2:
            QMessageBox.warning(self, "Import", "A mixture needs at least 2 components.")
            return
        if len(f_idx) > len(LETTERS):
            QMessageBox.warning(self, "Import", "Too many factors.")
            return
        n = len(self.raw)
        factors, fvals = [], np.full((n, len(f_idx)), np.nan)
        for k, j in enumerate(f_idx):
            c = self.cols[j]
            vals = [r[j] for r in self.raw]
            if c["role"] == "categoric":
                levels = list(dict.fromkeys(v for v in vals if v != ""))
                if len(levels) < 2:
                    QMessageBox.warning(self, "Import", f"Categoric factor {c['name']} needs at least 2 levels.")
                    return
                if len(levels) > 20:
                    QMessageBox.warning(self, "Import", f"{c['name']} has {len(levels)} levels - should this column "
                                                         "be numeric or ignored?")
                    return
                for i, v in enumerate(vals):
                    if v != "":
                        fvals[i, k] = levels.index(v)
                factors.append(Factor(c["name"], c["unit"], 0, 0, kind="categoric", levels=levels))
                continue
            nums = [_num(v) if v != "" else math.nan for v in vals]
            if any(v is None for v in nums):
                bad = next(i for i, v in enumerate(nums) if v is None)
                QMessageBox.warning(self, "Import", f"Column {c['name']}, row {bad + 1} is not a number: '{vals[bad]}'.")
                return
            arr = np.array(nums, float)
            fvals[:, k] = arr
            ok = arr[np.isfinite(arr)]
            if len(ok) == 0 or ok.max() == ok.min():
                QMessageBox.warning(self, "Import", f"Factor {c['name']} is constant / empty - mark it as Ignore.")
                return
            if c["role"] == "component":
                factors.append(Factor(c["name"], c["unit"], float(ok.min()), 0.0, upper=float(ok.max())))
            else:
                factors.append(Factor(c["name"], c["unit"], float(ok.min()), float(ok.max())))
        responses, yvals = [], np.full((n, len(r_idx)), np.nan)
        for k, j in enumerate(r_idx):
            c = self.cols[j]
            for i, r in enumerate(self.raw):
                if r[j] == "":
                    continue
                v = _num(r[j])
                if v is None:
                    QMessageBox.warning(self, "Import", f"Response {c['name']}, row {i + 1} is not a number: '{r[j]}'.")
                    return
                yvals[i, k] = v
            responses.append(Response(c["name"], c["unit"]))
        total = None
        if comp:
            total = self.sp_total.value()
            sums = np.nansum(fvals, axis=1)
            off = np.abs(sums - total) > 1e-6 * max(1, total) + 1e-9
            if off.any():
                if QMessageBox.question(self, "Import", f"In {int(off.sum())} rows the components do not sum to the "
                                                        f"total {total:g}. Continue (those rows are still analyzed)?") \
                        != QMessageBox.Yes:
                    return
            if sum(f.low for f in factors) >= total:
                QMessageBox.warning(self, "Import", "Sum of component minimums ≥ total - check the mixture total.")
                return
        try:
            self.project = Project.from_table(factors, responses, fvals, yvals, total)
        except (ValueError, np.linalg.LinAlgError) as exc:
            QMessageBox.warning(self, "Import", str(exc))
            return
        self.accept()
